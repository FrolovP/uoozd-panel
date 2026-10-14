# ДЕМО-РЕЖИМ ПАНЕЛИ РЕЕСТРА — для показа. Все суммы, предметы и участники вымышленные.
#
# Положите рядом: reestr_server.py, этот файл и папку «Документы для показа».
# Запуск:  python демо.py            — собрать демо-реестр заново (от сегодняшней даты) и открыть панель
#          python демо.py --оставить  — открыть панель на уже собранном демо-реестре
#
# Боевой реестр и рабочие папки не используются: всё лежит в папке «Демо-реестр» рядом с файлом.

import random
import sys as _sys
for _st in (_sys.stdout, _sys.stderr):   # консоль Windows не умеет значки ✅ ℹ️ — не падаем, а заменяем
    try:
        _st.reconfigure(encoding='utf-8', errors='replace')
    except (AttributeError, ValueError):
        pass
_missing = []
try:
    import pypdf  # noqa: F401  — нужна только для PDF; без неё демо работает, PDF не читаются
except ImportError:
    print('Внимание: не установлена библиотека pypdf — PDF-документы читаться не будут, остальное работает.')
    print('Установить:  python -m pip install pypdf')
for _m, _pkg in (('openpyxl', 'openpyxl'), ('docx', 'python-docx')):
    try:
        __import__(_m)
    except ImportError:
        _missing.append(_pkg)
if _missing:
    print('Не установлены библиотеки:', ', '.join(_missing))
    print('Откройте командную строку и выполните:  python -m pip install ' + ' '.join(_missing))
    input('Нажмите Enter, чтобы закрыть окно…')
    _sys.exit(1)
from openpyxl import load_workbook
import shutil
import sys
from datetime import date, datetime, timedelta
from pathlib import Path

import reestr_server as r

HERE = Path(__file__).resolve().parent
BASE = HERE / 'Демо-реестр'
DOCS = HERE / 'Документы для показа'

Y0 = date.today().year                     # текущий год; прошлый — архив, следующий — уже принимает извещения
r.ПАПКА_РЕЕСТРА = BASE
r.ГОДЫ = {
    Y0 - 1: {'реестр': BASE / str(Y0 - 1) / f'Реестр {Y0 - 1}.xlsx', 'папка': BASE / str(Y0 - 1), 'приём': False},
    Y0: {'реестр': BASE / f'Реестр {Y0}.xlsx', 'папка': BASE},
    Y0 + 1: {'реестр': BASE / str(Y0 + 1) / f'Реестр {Y0 + 1}.xlsx', 'папка': BASE / str(Y0 + 1)},
}
r.ПАПКА_ПАНЕЛИ = BASE / 'Панель'
r.ПАПКА_КАТАЛОГА = r.ПАПКА_ПАНЕЛИ / 'Каталог'
r.УВЕДОМЛЕНИЯ_ДЕМО = True                 # журнал уведомлений: «отправлено (демо)», без реальной отправки
r.ОПАиЗ_НАЧАЛЬНИК = 'Сальникова Е.А.'
r.ВХОД_ПО_ПАРОЛЮ = False                  # в демо без входа: роль выбирается в шапке
REG = lambda y=Y0: r.year_obj(y).registry
r.ФАЙЛ_АНАЛИТИКОВ = BASE / 'Реестр аналитиков.xlsx'
r.ПОРТ = 8090
r.ИНТЕРВАЛ = 10
GONCHAROVA = 'Гончарова К.А.'
r.ОТВЕТСТВЕННЫЕ = [(GONCHAROVA, 'goncharova_k_a@staff.sechenov.ru'),
                   ('Кузнецова М.А.', 'kuznetsova_m_a@staff.sechenov.ru'),
                   ('Кузнецова А.С.', 'kuznetsova_a_s@staff.sechenov.ru'),
                   ('Терешкин А.А.', 'tereshkin_a_a@staff.sechenov.ru'),
                   ('Миранова В.С.', '')]
PEOPLE = [x[0] for x in r.ОТВЕТСТВЕННЫЕ][:4]
DEMO_USERS = [  # в демо пароли не нужны — роль выбирается в шапке
    {'login': 'miranova', 'fio': 'Миранова В.С.', 'dept': 'ОАО', 'admin': True, 'can_delete': True},
    {'login': 'goncharova', 'fio': 'Гончарова К.А.', 'dept': 'ОАО', 'admin': False, 'can_delete': False},
    {'login': 'kuznetsova_a_s', 'fio': 'Кузнецова А.С.', 'dept': 'ОАО', 'admin': False, 'can_delete': False},
    {'login': 'nedorosleva', 'fio': 'Недорослева Н.П.', 'dept': 'ОПАиЗ', 'admin': False, 'can_delete': False},
    {'login': 'ovrk', 'fio': 'Сотрудник ОВРК', 'dept': 'ОВРК', 'admin': False, 'can_delete': False},
]

# ---------- «часы» генератора: записи получают правдоподобные даты ----------
NOW = [datetime.now()]


class DemoDateTime(datetime):
    @classmethod
    def now(cls, tz=None):
        return NOW[0]


def at(d, hh=10, mm=0):
    NOW[0] = datetime(d.year, d.month, d.day, hh, mm, random.randint(0, 59))


ru = lambda d: d.strftime('%d.%m.%Y')
iso = lambda d: d.strftime('%Y-%m-%d')

COMPANIES = [('ООО "АЛЬФА"', '0000000001'), ('ООО "БЕТА"', '0000000002'), ('ООО "ГАММА"', '0000000003'),
             ('ООО "ДЕЛЬТА"', '0000000004'), ('ООО "ЭПСИЛОН"', '0000000005'), ('ООО "ДЗЕТА"', '0000000006'),
             ('ООО "ТЕТА"', '0000000008'), ('ООО "ЙОТА"', '0000000009'), ('ООО "КАППА"', '0000000010'),
             ('ООО "СИГМА"', '0000000011'), ('ООО "ОМЕГА"', '0000000012'),
             ('Индивидуальный предприниматель ПЕТРОВ ПЁТР ПЕТРОВИЧ', '000000000013'),
             ('Индивидуальный предприниматель СИДОРОВА АННА СЕРГЕЕВНА', '000000000014')]
SUBJ_K = ['Поставка медицинского оборудования', 'Поставка лабораторного оборудования', 'Выполнение работ по текущему ремонту помещений',
          'Выполнение работ по капитальному ремонту', 'Оказание услуг по техническому обслуживанию оборудования',
          'Поставка компьютерной техники', 'Поставка лекарственных препаратов', 'Поставка медицинских изделий']
SUBJ_NK = ['Поставка расходных материалов', 'Поставка канцелярских товаров', 'Поставка учебного оборудования', 'Поставка мебели',
           'Поставка хозяйственных товаров', 'Поставка медицинских изделий', 'Поставка реагентов',
           'Оказание услуг по обслуживанию кондиционеров', 'Поставка оргтехники', 'Поставка спецодежды']
TAIL = ' для нужд ФГАОУ ВО Первый МГМУ им. И.М. Сеченова Минздрава России (Сеченовский Университет) в 2026 году'
REJECT = 'Заявка не соответствует требованиям п. 3.2 документации (не представлены документы)'


def submit(form, who):
    form = dict(form, author=who.split()[0])
    code, ans = r.submit_from_panel(r.form_to_d(form), '')
    if code not in (200, 202):
        print('  ⚠', ans['message'])


def apply(action, number, payload, who, token=''):
    data = {'action': action, 'number': number, action: payload, 'author': who.split()[0], 'token': token}
    code, ans = r.handle_apply(data)
    if code not in (200, 202):
        print('  ⚠', ans.get('message'))


def proto_doc(kind, number, p, who):
    """Сформировать протокол, как это делает сотрудник, и приложить его к записи."""
    try:
        info = r.protocol_info(number)
        content = r.build_protocol_docx(kind, info, r._clean_protocol(p), {}, r.responsible_short(who, 'k'))
        return r.keep_upload(f"Протокол {number}-{p['seq']} {info['oao']}.docx", content)
    except ValueError:
        return ''


# ---------- шесть сценариев по реальным документам (обезличены) ----------

def load_doc(name):
    f = next(DOCS.rglob(name + '*'), None)
    if not f:
        print('  нет файла', name)
        return None
    import base64
    return r.handle_doc({'b64': base64.b64encode(f.read_bytes()).decode(), 'filename': f.name})


def doc_notice(prefix, who, when):
    res = load_doc(prefix)
    if not res:
        return
    at(when)
    form = res['form']
    form['responsible'] = who
    form['author'] = who.split()[0]
    code, ans = r.submit_from_panel(r.form_to_d(form), res.get('body', ''))


def doc_apply(prefix, who, when, contract_date=None):
    res = load_doc(prefix)
    if not res:
        return
    at(when, 15)
    if res['doc'] == 'protocol':
        apply('protocol', res['match']['number'], res['protocol'], who, res['token'])
    else:
        c = res['contract']
        c['date'] = iso(contract_date or when)
        apply('contract', res['match']['number'], c, who, res['token'])


def seed_documents():
    d = lambda s: datetime.strptime(s, '%d.%m.%Y').date()
    G, T = GONCHAROVA, 'Терешкин А.А.'
    doc_notice('1388зэк — извещение', T, d('08.07.2026'))
    doc_notice('1535зэа — извещение', G, d('16.07.2026'))
    doc_notice('415р — извещение', G, d('31.07.2026'))
    doc_notice('1700зэа — извещение', G, d('31.08.2026'))
    doc_notice('1718зцп — извещение', G, d('31.08.2026'))
    doc_notice('542зцп — извещение', G, d('03.09.2026'))
    # Внесено заранее; остальное — загрузить вживую на показе (см. README).
    doc_apply('1388зэк — протокол 1', T, d('24.07.2026'))
    doc_apply('1388зэк — протокол 2', T, d('28.07.2026'))
    doc_apply('1535зэа — протокол 1', G, d('05.08.2026'))
    doc_apply('415р — протокол 1', G, d('11.08.2026'))
    doc_apply('1700зэа — протокол 1', G, d('16.09.2026'))
    doc_apply('1700зэа — договор', G, d('29.09.2026'))
    doc_apply('542зцп — протокол 1', G, d('10.09.2026'))
    doc_apply('542зцп — договор', G, d('24.09.2026'))


# ---------- вымышленные закупки во всех состояниях ----------

def synthetic(today, count=52):
    rnd = random.Random(2026)
    used_oao = set()
    plan = []
    # Обязательные «сегодняшние» события у Гончаровой — чтобы «Мой день» не был пустым.
    forced = [('auc', 0, 'bid'), ('auc', 0, 'review'), ('nk', 0, 'result'), ('konk', 0, 'review'), ('nk', 0, 'bid'),
              ('auc', -3, 'review'), ('nk', -4, 'result'), ('auc', 2, 'trade'), ('nk', 3, 'result')]
    for i in range(count):
        if i < len(forced):
            kind, off, stage = forced[i]
            who = GONCHAROVA
        else:
            kind = rnd.choices(['auc', 'konk', 'nk'], [40, 12, 48])[0]
            off, stage = rnd.randint(-75, 12), 'bid'
            who = rnd.choice(PEOPLE)
        # Сдвигаем всю цепочку дат так, чтобы нужный этап пришёлся на today+off.
        gaps = {'auc': dict(bid=0, review=2, trade=3, result=4), 'konk': dict(bid=0, review=4, result=5),
                'nk': dict(bid=0, result=1)}[kind]
        bid = today + timedelta(days=off - gaps.get(stage, 0))
        while bid.weekday() >= 5:
            bid += timedelta(days=1)
        pub = bid - timedelta(days=rnd.randint(7, 16))
        dates = {k: bid + timedelta(days=v) for k, v in gaps.items()}
        while True:
            n = rnd.randint(1600, 2400) if kind != 'nk' else rnd.randint(300, 1300)
            oao = f"{n}{'зэа' if kind == 'auc' else 'зэк' if kind == 'konk' else rnd.choice(['зцп', 'зцп', 'зцп', 'р'])}"
            if oao not in used_oao:
                used_oao.add(oao)
                break
        number = (f"326163{rnd.randint(10000, 99999)}" if kind != 'nk'
                  else f"COM{pub.strftime('%d%m%y')}00{rnd.randint(100, 999)}")
        subj = rnd.choice(SUBJ_NK if kind == 'nk' else SUBJ_K)
        smp = kind == 'nk' and rnd.random() < .5
        p2030 = rnd.random() < .15
        name = (subj + f' (лот № {rnd.randint(1, 40)})' + (' в рамках программы «Приоритет-2030»' if p2030 else '') + TAIL +
                (' (для субъектов малого и среднего предпринимательства)' if smp else ''))
        nmck = round(rnd.choice([rnd.uniform(150_000, 990_000), rnd.uniform(1e6, 6e6), rnd.uniform(6e6, 25e6)]), 2) if kind != 'nk' \
            else round(rnd.uniform(80_000, 3_000_000), 2)
        plan.append(dict(kind=kind, oao=oao, number=number, name=name, nmck=nmck, who=who, pub=pub, dates=dates, smp=smp,
                         p2030=p2030, rnd=random.Random(sum(map(ord, oao)) * 7919)))
    return plan


def run_synthetic(today, plan):
    events = []
    for x in plan:
        events.append((x['pub'], 0, 'notice', x))
    events.sort(key=lambda e: (e[0], e[1]))
    for when, _, _, x in events:
        if when > today:
            continue
        at(when, 9 + x['rnd'].randint(0, 7), x['rnd'].randint(0, 59))
        submit(notice_form(x), x['who'])
    # Протоколы и договоры — по датам этапов.
    later = []
    for x in plan:
        later += outcomes(x, today)
    later.sort(key=lambda e: e[0])
    for d, action, x, payload, kind_key in later:
        at(d, 11 + x['rnd'].randint(0, 6), x['rnd'].randint(0, 59))
        token = proto_doc(kind_key, x['number'], payload, x['who']) if action == 'protocol' and kind_key != 'konk_result' else ''
        apply(action, x['number'], payload, x['who'], token)
    return plan


def notice_form(x):
    if True:
        k, dts = x['kind'], x['dates']
        form = {'kind': 'nk' if k == 'nk' else 'k', 'oao': x['oao'], 'number': x['number'], 'name': x['name'],
                'price': f"{x['nmck']:.2f}", 'publish': iso(x['pub']), 'bid': iso(dts['bid']),
                'review': iso(dts['review']) if 'review' in dts else '', 'trade': iso(dts['trade']) if 'trade' in dts else '',
                'result': iso(dts['result']), 'responsible': x['who'], 'letter': x.get('letter') or x['rnd'].choice(['Р', 'К', 'мед', 'Л', 'ТО']),
                'one_c': x['rnd'].choice(['Да', 'Нет']), 'subordination': x['rnd'].choice(['ВУЗ', 'КЦ']),
                'subsidy': '2030' if x['p2030'] else x['rnd'].choice(['', '', '', 'Соглашение']),
                'site_url': ('https://com.roseltorg.ru' if k == 'nk' else x['rnd'].choice(['https://www.astgoz.ru', 'https://etpgpb.ru'])),
                'smp': x['smp'], 'p2030': x['p2030']}
    return form


def outcomes(x, today, p_done=.82, p_contract=.8):
    """Протоколы и договор закупки: [(дата, 'protocol'|'contract', x, данные, вид протокола)]."""
    later = []
    if True:
        rnd, k, dts = x['rnd'], x['kind'], x['dates']
        n_apps = rnd.choices([0, 1, 2, 3, 4], [6, 30, 34, 20, 10])[0]
        names = rnd.sample(COMPANIES, n_apps)
        parts = []
        for i, (nm, inn) in enumerate(names, 1):
            rej = n_apps > 1 and rnd.random() < .15
            parts.append({'id': str(rnd.randint(10000, 99999)) if k != 'nk' else str(i),
                          'date': ru(dts['bid'] - timedelta(days=rnd.randint(1, 5))) + f" {rnd.randint(8, 18):02d}:{rnd.randint(0, 59):02d}",
                          'name': nm, 'inn': inn, 'kpp': '', 'decision': 'Отклонена' if rej else 'Допущена',
                          'reason': REJECT if rej else '-', 'smp': rnd.choice(['микро', 'малое', 'малое', 'среднее', '', '', '', 'СОНКО'])})
        adm = [p for p in parts if p['decision'] == 'Допущена']
        for p in adm:
            p['price'] = round(x['nmck'] * rnd.uniform(.86, 1.0) if len(adm) > 1 else x['nmck'], 2)
        win = min(adm, key=lambda p: p['price']) if adm else None
        done = lambda d: d <= today and rnd.random() < p_done
        stages = []
        if k == 'konk':
            stages = [('konk_open', 1, dts['bid']), ('konk_review', 2, dts['review'])]
            final = ('konk_result', 3, dts['result'])
        elif k == 'auc':
            if len(adm) <= 1:
                final = ('auc_single', 1, dts['review'])
            else:
                stages = [('auc_review', 1, dts['review'])]
                final = ('auc_result', 2, dts['result'])
        else:
            final = ('nk_result', 1, dts['result'])
        ok = True
        for kind_key, seq, d in stages + [final]:
            if not ok or not done(d):
                break
            is_final = (kind_key, seq, d) == final
            p = {'seq': seq, 'date': ru(d), 'type': r.KIND_TO_TYPE[kind_key], 'oao': x['oao'], 'number': x['number'],
                 'form': 'конкурс' if k == 'konk' else 'аукцион' if k == 'auc' else '',
                 'apps': n_apps, 'participants': [dict(p, price=p.get('price') if is_final else None) for p in parts],
                 'admitted': len(adm), 'rejected': n_apps - len(adm),
                 'bidders': len(adm) if kind_key == 'auc_result' else 0, 'failed': len(adm) == 1}
            if is_final and win:
                p['winner'] = {'name': win['name'], 'inn': win['inn'], 'id': win['id'], 'price': win['price'], 'smp': win['smp']}
            later.append((d, 'protocol', x, p, kind_key))
            if is_final and win:
                cd = d + timedelta(days=rnd.randint(9, 22))
                while x.get('weekday_only', True) and cd.weekday() >= 5 and rnd.random() < .85:
                    cd += timedelta(days=1)
                if cd <= today and rnd.random() < p_contract:
                    price = round(win['price'] * (rnd.choice([1, 1, 1, .98, .995])), 2)
                    later.append((cd, 'contract', x, {'contract_no': x['number'], 'date': ru(cd), 'price': price,
                                                      'supplier': win['name'], 'inn': win['inn']}, None))
    return later


LETTER_W = [('Р', 22), ('мед', 20), ('К', 14), ('Л', 12), ('ТО', 12), ('ремонт обр', 6), ('Уборка', 3), ('охрана', 2),
            ('питание Звенигород', 2), ('Хлеб', 2), ('снег', 1), ('Гардероб', 1), ('ТО обр', 3)]
MONTH_W = [5, 8, 10, 11, 9, 10, 7, 8, 12, 12, 11, 6]   # сезонность публикаций


def history(today, count=240, year=None, seed=1717, used=None):
    """Закупки года — для аналитики. Пишутся пачками (по неделе), без файлов протоколов.
    year=None — текущий год до «сегодня − 80 дней»; иначе — весь указанный год."""
    rnd = random.Random(seed)
    year = year or today.year
    if year == today.year:
        start, end = date(year, 1, 12), today - timedelta(days=80)
    else:
        start, end = date(year, 1, 12), date(year, 12, 18)
        today = date(year + 1, 1, 31)       # всё, что было в том году, уже прошло
    days = [start + timedelta(days=i) for i in range((end - start).days + 1)]
    weights = [MONTH_W[d.month - 1] * (0 if d.weekday() >= 5 else 1) for d in days]
    used = used if used is not None else set()
    used |= {'1388зэк', '1535зэа', '1700зэа', '415р', '542зцп', '1718зцп', '2001зцп'}
    plan = []
    for i in range(count):
        kind = rnd.choices(['auc', 'konk', 'nk'], [40, 10, 50])[0]
        bid = rnd.choices(days, weights)[0]
        gaps = {'auc': dict(bid=0, review=2, trade=3, result=4), 'konk': dict(bid=0, review=4, result=5),
                'nk': dict(bid=0, result=1)}[kind]
        pub = bid - timedelta(days=rnd.randint(7, 16))
        if pub.year < year:
            pub = date(year, 1, 9)
        while True:
            n = rnd.randint(20, 1590) if kind != 'nk' else rnd.randint(1301, 2999)
            oao = f"{n}{'зэа' if kind == 'auc' else 'зэк' if kind == 'konk' else rnd.choice(['зцп', 'зцп', 'зцп', 'зцп', 'р'])}"
            if oao not in used:
                used.add(oao)
                break
        number = (f"326{rnd.randint(10000000, 16000000)}" if kind != 'nk' else f"COM{pub.strftime('%d%m%y')}0{rnd.randint(1000, 9999)}")
        subj = rnd.choice(SUBJ_NK if kind == 'nk' else SUBJ_K)
        smp = kind == 'nk' and rnd.random() < .5
        p2030 = rnd.random() < .15
        name = (subj + f' (лот № {rnd.randint(1, 40)})' + (' в рамках программы «Приоритет-2030»' if p2030 else '') + TAIL.replace('2026', str(year)) +
                (' (для субъектов малого и среднего предпринимательства)' if smp else ''))
        nmck = round(rnd.choice([rnd.uniform(150_000, 990_000), rnd.uniform(1e6, 6e6), rnd.uniform(6e6, 25e6)]), 2) if kind != 'nk' \
            else round(rnd.uniform(80_000, 3_000_000), 2)
        who = rnd.choices(PEOPLE, [34, 24, 22, 20])[0]
        plan.append(dict(kind=kind, oao=oao, number=number, name=name, nmck=nmck, who=who, pub=pub,
                         dates={k: bid + timedelta(days=v) for k, v in gaps.items()}, smp=smp, p2030=p2030,
                         letter=rnd.choices([x[0] for x in LETTER_W], [x[1] for x in LETTER_W])[0],
                         rnd=random.Random(sum(map(ord, oao)) * 7919 + i)))
    # извещения — пачкой за неделю
    plan.sort(key=lambda x: x['pub'])
    weeks = {}
    for x in plan:
        weeks.setdefault(x['pub'] - timedelta(days=x['pub'].weekday()), []).append(x)
    for wk in sorted(weeks):
        at(weeks[wk][0]['pub'], 10, 5)
        r.write_entries([r.form_to_d(notice_form(x)) for x in weeks[wk]], REG(year))
    # протоколы и договоры — сразу в книгу, одной записью
    later = []
    for x in plan:
        later += outcomes(x, today, p_done=.97, p_contract=.93)
    later.sort(key=lambda e: e[0])
    wb = load_workbook(REG(year))
    for d, action, x, payload, kind_key in later:
        loc = r.locate(wb, x['number'])
        if not loc:
            continue
        up = r.protocol_updates(r._clean_protocol(payload), loc) if action == 'protocol' else r.contract_updates(r._clean_contract(payload), loc)
        for ch in r.make_changes(loc, up):
            if ch['state'] != 'formula':
                loc['ws'].cell(loc['row'], ch['col']).value = r.safe(ch['value']) if ch['value'] != '' else None
    r.save_workbook(wb, REG(year))
    return plan


def corrections(today, plan):
    """Пара исправлений через карточку — чтобы была видна история изменений."""
    done = 0
    for x in plan:
        if done >= 3:
            break
        loc = r.open_location(x['number'])
        if not loc or not loc['cols'].get('winner'):
            continue
        cur = loc['ws'].cell(loc['row'], loc['cols']['winner']).value
        if not cur or 'НЕ СОСТОЯЛ' in str(cur):
            continue
        at(today - timedelta(days=done + 1), 16, 20)
        wrong = str(cur).replace('ООО', 'ОО0') if done == 0 else str(cur) + ' '
        col = loc['cols']['winner']
        r.handle_apply({'action': 'edit', 'number': x['number'], 'values': {str(col): wrong}, 'author': x['who'].split()[0]})
        at(today - timedelta(days=done), 9, 40)
        r.handle_apply({'action': 'edit', 'number': x['number'], 'values': {str(col): str(cur)}, 'author': 'Гончарова'})
        done += 1


def legacy_names(plan):
    """Несколько строк «как было раньше» (КузнецоваМ, полное ФИО, без инициалов) —
    чтобы на показе было что привести к виду «Фамилия И.О.» кнопкой «Инициалы»."""
    old = {'Кузнецова М.А.': ['КузнецоваМ', 'Кузнецова Мария Александровна'],
           'Терешкин А.А.': ['Терешкин'], 'Кузнецова А.С.': []}
    wb = load_workbook(REG())
    n = 0
    for x in plan:
        if x['who'] not in old or not old[x['who']]:
            continue
        loc = r.locate(wb, x['number'])
        if loc and loc['cols'].get('responsible'):
            loc['ws'].cell(loc['row'], loc['cols']['responsible']).value = old[x['who']].pop(0)
            n += 1
    r.save_workbook(wb, REG())


def report_history(today):
    """Прошлые недельные отчёты «отправлены» в пятницу в 17:00; пара договоров внесена уже после отправки —
    они видны в текущем отчёте с пометкой."""
    import json
    r.cache['mtime'] = None
    r.refresh_cache()                       # первая запись «когда появились»: всё внесено вовремя
    rows = r.cache['rows']
    at(today, 9, 0)
    cur = r.week_of_moment(datetime(today.year, today.month, today.day, 9, 0))
    prev = cur - timedelta(days=7)
    path = r.ПАПКА_ПАНЕЛИ / r.ФАЙЛ_ПОЯВЛЕНИЙ
    seen = json.loads(path.read_text(encoding='utf-8'))
    late = sorted([x for x in rows if x.get('contract_date') and prev - timedelta(days=4) <= date.fromisoformat(x['contract_date']) <= prev],
                  key=lambda x: x['contract_date'], reverse=True)[:2]
    for x, t in zip(late, ['18:40', '11:15']):
        moment = datetime.combine(prev if t == '18:40' else prev + timedelta(days=1), datetime.strptime(t, '%H:%M').time())
        seen['c:' + r.item_id(x)] = moment.isoformat(timespec='seconds')
    path.write_text(json.dumps(seen, ensure_ascii=False), encoding='utf-8')
    rnd = random.Random(55)
    for i in range(12, 0, -1):
        day = cur - timedelta(days=7 * i)
        if day.year < today.year:
            continue
        at(day, 16, 40)
        r.save_ep(day, {'count': str(rnd.randint(18, 70)), 'sum': f"{rnd.uniform(2e6, 9e7):.2f}".replace('.', ',')})
        at(day, 17, 0)
        r.fix_week(day, 'панель', auto=True)
    at(today, 9, 0)
    r.save_ep(cur, {'count': '', 'sum': ''})


def next_year(today, used):
    """Закупки на следующий год: опубликованы в последние недели, вскрытие — впереди."""
    rnd = random.Random(2027)
    plan = []
    for i in range(7):
        kind = rnd.choice(['auc', 'auc', 'nk', 'konk', 'nk'])
        pub = today - timedelta(days=rnd.randint(1, 18))
        while pub.weekday() >= 5:
            pub -= timedelta(days=1)
        bid = today + timedelta(days=rnd.randint(4, 30))
        while bid.weekday() >= 5:
            bid += timedelta(days=1)
        gaps = {'auc': dict(bid=0, review=2, trade=3, result=4), 'konk': dict(bid=0, review=4, result=5),
                'nk': dict(bid=0, result=1)}[kind]
        while True:
            n = rnd.randint(2400, 2700) if kind != 'nk' else rnd.randint(1400, 1600)
            oao = f"{n}{'зэа' if kind == 'auc' else 'зэк' if kind == 'konk' else 'зцп'}"
            if oao not in used:
                used.add(oao)
                break
        subj = rnd.choice(SUBJ_NK if kind == 'nk' else SUBJ_K)
        x = dict(kind=kind, oao=oao, name=subj + f' (лот № {rnd.randint(1, 40)})' + TAIL.replace(str(today.year), str(today.year + 1)),
                 number=(f"326{rnd.randint(16000000, 16999999)}" if kind != 'nk' else f"COM{pub.strftime('%d%m%y')}0{rnd.randint(1000, 9999)}"),
                 nmck=round(rnd.uniform(300_000, 9e6), 2), who=rnd.choice(PEOPLE), pub=pub, smp=kind == 'nk', p2030=False,
                 dates={k: bid + timedelta(days=v) for k, v in gaps.items()}, letter=rnd.choice(['Р', 'мед', 'К', 'ТО']),
                 rnd=random.Random(i * 31 + 7))
        at(pub, 11, 10)
        form = dict(notice_form(x), year=str(today.year + 1))
        submit(form, x['who'])
        plan.append(x)
    return plan


def protocol_folder(today, plan):
    """Кладём в «Протоколы\\входящие» три документа: один внесётся сам, два попадут на проверку."""
    import json
    folder = r.year_obj(Y0).protocols_in()
    folder.mkdir(parents=True, exist_ok=True)
    r.cache['mtime'] = None
    r.refresh_cache(force=True)
    rows = {x['number']: x for x in r.cache['rows']}
    waiting = [x for x in plan if x['kind'] == 'nk' and rows.get(x['number'], {}).get('status') in ('overdue', 'review')
               and not rows[x['number']].get('winner')]
    done = [x for x in plan if x['kind'] == 'nk' and rows.get(x['number'], {}).get('winner')
            and 'НЕ СОСТОЯЛ' not in rows[x['number']]['winner']]

    def proto(x, price_k, name, inn, info=None):
        info = info or r.protocol_info(x['number'])
        price = round(x['nmck'] * price_k, 2)
        p = {'seq': 1, 'date': ru(x['dates']['result']), 'type': 'review_result', 'oao': x['oao'], 'number': x['number'],
             'apps': 2, 'admitted': 2, 'rejected': 0, 'bidders': 0, 'failed': False,
             'participants': [{'id': '1', 'date': ru(x['dates']['bid'] - timedelta(days=2)) + ' 10:15', 'name': name, 'inn': inn,
                               'decision': 'Допущена', 'reason': '-', 'smp': 'малое', 'price': price},
                              {'id': '2', 'date': ru(x['dates']['bid'] - timedelta(days=1)) + ' 15:40', 'name': 'ООО "ОМЕГА"',
                               'inn': '0000000012', 'decision': 'Допущена', 'reason': '-', 'smp': '', 'price': round(price * 1.04, 2)}],
             'winner': {'name': name, 'inn': inn, 'id': '1', 'price': price, 'smp': 'малое'}}
        return r.build_protocol_docx('nk_result', info, r._clean_protocol(p), {}, r.responsible_short(x['who'], 'k'))

    if waiting:
        x = waiting[0]
        (folder / f"{x['oao']} — протокол рассмотрения и итогов.docx").write_bytes(proto(x, .93, 'ООО "СИГМА"', '0000000011'))
    if done:
        x = done[0]
        (folder / f"{x['oao']} — протокол (повторно, другая цена).docx").write_bytes(proto(x, .81, 'ООО "КАППА"', '0000000010'))
    if waiting:
        info = dict(r.protocol_info(waiting[-1]['number']), number='COM0101990009999', oao='9999зцп')
        fake = dict(waiting[-1], number='COM0101990009999', oao='9999зцп')
        (folder / 'протокол по закупке, которой нет в реестре.docx').write_bytes(proto(fake, .9, 'ООО "БЕТА"', '0000000002', info))


def analysts_register(today):
    """Реестр аналитиков (Excel делопроизводителя) с датами «на публикацию»: часть уже опубликована,
    часть ждёт, одна просрочена."""
    from openpyxl import Workbook
    r.cache['mtime'] = None
    r.refresh_cache(force=True)
    heads = ['Входящий номер ОАО', ' ', '№ КЦ', 'КОД', 'Подписант', 'подразделение ', '№ ПЭО', 'План-график', 'Примечания',
             'Название', 'Сумма', 'Вид закупки   (конкурс, аукцион)', 'В ПЭО', 'Ответственный в ОАО          ',
             'Дата поступления в ОАО ', 'Дата передачи в юр.отдел', 'в бухгалтерии', 'прорекор', 'на публикацию',
             'Подтверждение публикации']
    wb = Workbook()
    ws = wb.active
    ws.title = str(Y0)
    ws.append(heads)
    suf = {'зэа': 'за-оа', 'зэк': 'зк-оа', 'зцп': 'зцп', 'р': 'р'}
    kinds = {'зэа': 'Закрытый аукцион', 'зэк': 'Закрытый конкурс', 'зцп': 'запрос ценовых предложений', 'р': 'редукцион'}
    rnd = random.Random(99)
    recent = sorted([x for x in r.cache['rows'] if x.get('year') == Y0 and x.get('pub') and
                     (today - date.fromisoformat(x['pub'])).days <= 20], key=lambda x: x['pub'])[-8:]

    def back_workdays(d, n):
        while n:
            d -= timedelta(days=1)
            if d.weekday() < 5:
                n -= 1
        return d
    for x in recent:
        m = __import__('re').match(r'(\d+)(\D+)', x['oao'])
        req = back_workdays(date.fromisoformat(x['pub']), rnd.choice([1, 2, 2, 3]))
        ws.append([f" {m.group(1)} {suf.get(m.group(2), m.group(2))} {req:%d.%m.%y}", '223', '', '310', '', 'ВУЗ', str(rnd.randint(2000, 2400)),
                   MONTHS_RU[req.month - 1], '', x['name'][:90], x['nmck'], kinds.get(m.group(2), ''), '', x['responsible'].split()[0],
                   '', '', '', '', req.strftime('%d.%m.%Y'), ''])
    pend = [(0, 'Гончарова К.А.', 'зэа'), (1, 'Кузнецова А.С.', 'зцп'), (1, 'Гончарова', 'зэк'), (4, 'Терешкин А.А.', 'зэа'),
            (0, 'Кузнецова М.А.', 'зцп'), (2, 'Гончарова К.А.', 'зцп')]
    for i, (ago, who, k) in enumerate(pend):
        req = back_workdays(today + timedelta(days=1), ago + 1) if ago else today
        n = 2500 + i * 7
        subj = rnd.choice(SUBJ_NK if k == 'зцп' else SUBJ_K)
        ws.append([f" {n} {suf[k]} {req:%d.%m.%y}", '223', '', '346', '', rnd.choice(['ВУЗ', 'КЦ']), str(rnd.randint(2400, 2600)),
                   MONTHS_RU[req.month - 1], '', subj + f' (лот № {rnd.randint(1, 30)})', round(rnd.uniform(3e5, 6e6), 2), kinds[k], '', who,
                   '', '', '', '', req.strftime('%d.%m.%Y'), ''])
    ws2 = wb.create_sheet(str(Y0 + 1))
    ws2.append(heads)
    ws2.append([f" 3 за-оа {today:%d.%m.%y}", '223', '', '310', '', 'КЦ', '2610', 'Январь', '', 'Поставка медицинского оборудования (лот № 1)',
                4_200_000, 'Закрытый аукцион', '', 'Терешкин', '', '', '', '', today.strftime('%d.%m.%Y'), ''])
    wb.save(r.ФАЙЛ_АНАЛИТИКОВ)


OZ_INIT_DEPT = {'Иванов И.И.': 'Лаборатория № 1', 'Петрова А.В.': 'Клинический центр', 'Сидоров П.П.': 'Институт биотехнологий',
                'Орлова Е.Н.': 'Научный центр', 'Белов С.А.': 'Кафедра фармакологии', 'Зайцева Т.М.': 'Учебный центр',
                'Морозов Д.К.': 'Институт биотехнологий', 'Лебедева Н.Ю.': 'Клинический центр'}
OZ_INIT = list(OZ_INIT_DEPT)
OZ_FIN = [('Внебюджет', ''), ('Грант', 'Приоритет'), ('Грант', 'РНФ'), ('Бюджет', 'Госзадание'), ('Софинансирование', 'ПИШ'), ('Целевая субсидия', 'Минздрав')]
OZ_KOSGU = ['310-осн.средства', '346-расх.материалы', '226-услуги/работы', '225-ремонт/поверка', '341-лек.средства', '320-ПО', '347-принадлежности']
OZ_TERMS = ['в течение 30 рабочих дней', 'в течение 60 календарных дней', 'в течение 2 месяцев', 'не позднее 15.12.{y}', 'в течение 45 календарных дней']
OZ_EP_SUBJ = ['Поставка реактивов', 'Поставка лекарственных препаратов', 'Поверка средств измерений', 'Обслуживание оборудования',
              'Поставка лабораторной посуды', 'Подписка на периодические издания', 'Услуги по дезинфекции', 'Поставка запасных частей']
OZ_PRORECTORS = ['Свистунов А.А.', 'Литвинова Т.М.', 'Бровко М.Ю.', 'Тарасов В.В.', 'Волкова О.С.']
OZ_PRIORITY = ['01.000.Е.129', 'П1.000.Б.01', '02.100.А.07', '01.200.Е.44', '03.000.Б.12']
# позиции по предмету: (наименование, ОКПД2, ед., цена, российское)
OZ_CATALOG = [
    ('медицинского оборудования', [('Монитор пациента прикроватный', '26.60.12.129', 'шт', 420000, True), ('Облучатель-рециркулятор бактерицидный', '32.50.50.190', 'шт', 38000, True),
                                   ('Электрокардиограф 12-канальный', '26.60.12.111', 'шт', 210000, True), ('Аппарат ИВЛ портативный', '32.50.21.121', 'шт', 1350000, False)]),
    ('лабораторного', [('Центрифуга лабораторная настольная', '28.99.39.190', 'шт', 185000, True), ('Шкаф ламинарный II класса', '28.25.14.110', 'шт', 640000, True),
                       ('Термостат суховоздушный', '28.93.15.190', 'шт', 96000, False), ('Весы аналитические', '28.29.31.110', 'шт', 154000, False)]),
    ('компьютерн', [('Системный блок', '26.20.15.000', 'шт', 78000, True), ('Монитор 27"', '26.20.17.110', 'шт', 26000, False), ('Ноутбук', '26.20.11.110', 'шт', 92000, True)]),
    ('оргтехник', [('МФУ лазерное', '26.20.18.000', 'шт', 48000, True), ('Принтер лазерный', '26.20.16.120', 'шт', 26000, False), ('Сканер документов', '26.20.16.150', 'шт', 41000, False)]),
    ('расходных', [('Перчатки нитриловые', '22.19.60.119', 'пар', 18, False), ('Пробирки вакуумные', '32.50.13.190', 'шт', 22, False), ('Наконечники для дозаторов', '22.29.29.190', 'уп', 640, False)]),
    ('реаг', [('Набор реагентов ИФА', '21.20.23.110', 'набор', 24000, False), ('Буферный раствор', '20.59.52.199', 'фл', 3800, False)]),
    ('лекарств', [('Натрия хлорид р-р 0,9%', '21.20.10.141', 'фл', 54, True), ('Лидокаин р-р', '21.20.10.193', 'амп', 38, True)]),
    ('медицинских изделий', [('Шприц инъекционный 5 мл', '32.50.13.110', 'шт', 6, True), ('Катетер периферический', '32.50.13.190', 'шт', 48, False), ('Бинт стерильный', '21.20.24.131', 'шт', 25, True)]),
    ('канцеляр', [('Бумага А4', '17.12.14.129', 'пач', 340, True), ('Ручка шариковая', '32.99.12.110', 'шт', 25, False)]),
    ('учебного', [('Тренажёр сердечно-лёгочной реанимации', '32.99.53.130', 'шт', 210000, True), ('Интерактивная панель', '26.40.33.110', 'шт', 310000, True)]),
    ('мебел', [('Стол лабораторный', '31.01.12.150', 'шт', 46000, True), ('Шкаф для документов', '31.01.11.150', 'шт', 21000, True), ('Кресло офисное', '31.01.11.110', 'шт', 9800, True)]),
    ('хозяйствен', [('Средство моющее', '20.41.32.119', 'л', 210, True), ('Мешки для мусора', '22.22.11.000', 'уп', 180, True)]),
    ('спецодежд', [('Халат медицинский', '14.12.30.131', 'шт', 1450, True), ('Костюм хирургический', '14.12.30.131', 'шт', 1900, True)]),
]


def _catalog(name):
    low = (name or '').lower()
    return next((items for key, items in OZ_CATALOG if key in low), None)


def _items_for(name, rnd, total=None):
    cat = _catalog(name)
    if not cat:
        return []
    out = []
    for nm, okpd, unit, price, rf in rnd.sample(cat, min(len(cat), rnd.randint(2, 3))):
        qty = rnd.choice([1, 2, 3, 5, 10]) if price > 5000 else rnd.choice([100, 200, 500, 1000])
        out.append({'name': nm, 'okpd': okpd, 'unit': unit, 'qty': qty, 'price': price, 'sum': round(qty * price, 2), 'rf': rf and rnd.random() < .85})
    if total:
        k = total / max(1, sum(x['sum'] for x in out))
        for x in out:
            x['price'] = round(x['price'] * k, 2)
            x['sum'] = round(x['qty'] * x['price'], 2)
    return out


def _short(name):
    return name.replace(TAIL, '').split(' (лот')[0].strip()


def opaiz_demo(today):
    """Реестр ОПАиЗ: СЗ → заявки (связаны с реестром ОАО) / договоры ЕП / счета, позиции, РФ, исполнение, проверка СМП."""
    import json as _json
    import re as _re
    rnd = random.Random(707)
    staff = r.staff_opaiz()
    who = lambda: {'fio': rnd.choice(staff)}
    r.ПАПКА_ПАНЕЛИ.joinpath(r.OZ_DIR).mkdir(parents=True, exist_ok=True)
    r.write_atomic(r._oz_path('инициаторы'), _json.dumps(OZ_INIT_DEPT, ensure_ascii=False, indent=1))
    lists = r.oz_lists()
    lists['org'] = sorted({c[0] for c in COMPANIES})
    lists['initiator'] = OZ_INIT[:]
    lists['dept'] = sorted(set(OZ_INIT_DEPT.values()))
    lists['priority'] = OZ_PRIORITY[:]
    r.oz_save_lists(lists)
    r.cache['mtime'] = None
    r.refresh_cache(force=True)
    rows = [x for x in r.cache['rows'] if x.get('year') == Y0 and x.get('pub') and x.get('oao') and _re.match(r'\d+', x['oao'])]
    seen, pick = set(), []
    for x in sorted(rows, key=lambda x: x['pub']):
        d = _re.match(r'\d+', x['oao']).group(0)
        if d not in seen:
            seen.add(d)
            pick.append(x)
    pick = sorted(rnd.sample(pick, min(55, len(pick))), key=lambda x: x['pub'])

    def new_sz(subject, din, inits, items, st='в работе', remarks=''):
        at(din, 9, 30)
        fin, prog = rnd.choice(OZ_FIN)
        num_ = rnd.randint(100, 999)
        row, _ = r.oz_save_row('sz', None, {
            'sz': f'№ {num_}/61-17 от {din:%d.%m.%Y}', 'din': din.isoformat(), 'author': inits[0], 'initiator': inits,
            'fin': fin, 'priority': rnd.choice(OZ_PRIORITY) if fin != 'Бюджет' else '', 'dept': OZ_INIT_DEPT[inits[0]],
            'prorector': rnd.choice(OZ_PRORECTORS), 'due': (din + timedelta(days=rnd.randint(14, 40))).isoformat(), 'resp': rnd.choice(staff),
            'subject': subject, 'remarks': remarks, 'fy': str(din.year), 'items': items, 'st': st,
            'done': (din + timedelta(days=rnd.randint(5, 20))).isoformat() if st == 'исполнена' else ''}, None, who())
        return row

    # 1) электронные заявки ← СЗ, связаны с реестром ОАО
    n = 0
    for x in pick:
        pub = date.fromisoformat(x['pub'])
        d1c = pub - timedelta(days=rnd.randint(6, 20))
        inits = rnd.sample(OZ_INIT, rnd.choice([1, 1, 2]))
        sz_items = _items_for(x['name'], rnd, x.get('nmck'))
        extra = _items_for(x['name'], rnd)[:1] if sz_items and rnd.random() < .35 else []     # часть потребности уйдёт в счёт
        for e in extra:
            e['name'] = e['name'] + ' (доп.)'
        sz = new_sz(_short(x['name']), d1c - timedelta(days=rnd.randint(5, 25)), inits, [dict(i) for i in sz_items + extra],
                    st='исполнена' if pub < today - timedelta(days=30) else 'в работе')
        at(d1c, 11)
        fin, prog = rnd.choice(OZ_FIN)
        m = _re.match(r'(\d+)', x['oao'])
        z_items = [dict(i, src=f"{sz['id']}#{k}") for k, i in enumerate(sz_items)]
        rf = 'да' if any(i['rf'] for i in z_items) and rnd.random() < .6 else 'нет'
        if rf == 'нет':
            for i in z_items:
                i['rf'] = False
        vals = {'status': 'sign', 'sz': sz['sz'].replace('№ ', ''), 'sz_links': [sz['id']], 'platform': rnd.choice(['АСТ ГОЗ', 'РосЭлТорг', 'ЭТП ГПБ']),
                'n1c': m.group(1), 'd1c': d1c.isoformat(), 'name': x['name'], 'nmck': x.get('nmck'), 'rf': rf,
                'smp': rnd.choice(['да', 'нет', 'нет']), 'kosgu': rnd.choice(OZ_KOSGU), 'fin': fin, 'program': prog, 'prorector': sz.get('prorector', ''),
                'priority': sz.get('priority', ''), 'resp': sz['resp'], 'initiator': inits, 'dept': OZ_INIT_DEPT[inits[0]],
                'period': rnd.choice(['Одномоментно', 'Партиями', 'По заявкам']), 'term': rnd.choice(OZ_TERMS).format(y=Y0), 'fy': str(Y0), 'items': z_items}
        row, _ = r.oz_save_row('zayavki', None, vals, None, who())
        n += 1
        if extra:     # остаток потребности СЗ — закупили по счёту
            dd = d1c + timedelta(days=rnd.randint(3, 15))
            if dd <= today:
                at(dd, 14)
                s = round(sum(e['sum'] for e in extra), 2)
                r.oz_save_row('invoices', None, {
                    'sz': sz['sz'], 'sz_links': [sz['id']], 'date_oao': dd.isoformat(), 'kosgu': vals['kosgu'], 'org': rnd.choice(COMPANIES)[0],
                    'inv': f'№{rnd.randint(1, 400)} от {dd - timedelta(days=2):%d.%m.%Y}', 'sum': s, 'subject': extra[0]['name'], 'rf': 'да' if extra[0]['rf'] else 'нет',
                    'terms': 'предоплата 100% в течение 5 рабочих дней', 'advance': '100%', 'paid': s if dd < today - timedelta(days=10) else '',
                    'paid_date': (dd + timedelta(days=4)).isoformat() if dd < today - timedelta(days=10) else '', 'resp': sz['resp'], 'dept_resp': inits,
                    'dept': OZ_INIT_DEPT[inits[0]], 'fin': fin, 'prorector': sz.get('prorector', ''), 'fy': str(Y0), 'clause': '8.1.7',
                    'items': [dict(e, src=f"{sz['id']}#{len(sz_items) + k}") for k, e in enumerate(extra)]}, None, who())
    # особые статусы
    data = r.oz_load('zayavki')['rows']
    for z, st, note in zip(rnd.sample(data, 3), ['ret', 'stop', 'repub'], ['возврат инициатору — нет обоснования НМЦК', 'приостановлено по письму проректора', 'перепубликация: изменено ТЗ']):
        r.oz_save_row('zayavki', z['id'], {'status': st, 'note': note}, None, who())
    # ещё не опубликованные: на подписании через ОАО и в работе у ОПАиЗ
    for i in range(4):
        at(today - timedelta(days=rnd.randint(1, 6)), 10)
        inits = rnd.sample(OZ_INIT, 1)
        subj = rnd.choice(SUBJ_NK + SUBJ_K)
        sz = new_sz(subj, today - timedelta(days=rnd.randint(12, 25)), inits, _items_for(subj, rnd))
        fin, prog = rnd.choice(OZ_FIN)
        r.oz_save_row('zayavki', None, {
            'status': 'sign', 'n1c': str(2500 + i * 7), 'd1c': (today - timedelta(days=rnd.randint(3, 16))).isoformat(), 'sz_links': [sz['id']], 'sz': sz['sz'].replace('№ ', ''),
            'name': subj + f' (лот № {rnd.randint(1, 30)})' + TAIL, 'nmck': round(rnd.uniform(3e5, 6e6), 2), 'items': [dict(x, src=f"{sz['id']}#{k}") for k, x in enumerate(sz.get('items') or [])],
            'method': rnd.choice(['ЗЦП', 'ЗЭА', 'ЗЭК']), 'kosgu': rnd.choice(OZ_KOSGU), 'fin': fin, 'program': prog, 'resp': sz['resp'],
            'initiator': inits, 'dept': OZ_INIT_DEPT[inits[0]], 'fy': str(Y0), 'term': rnd.choice(OZ_TERMS).format(y=Y0)}, None, who())
    for i in range(5):
        at(today - timedelta(days=rnd.randint(0, 3)), 15)
        inits = rnd.sample(OZ_INIT, 1)
        subj = rnd.choice(SUBJ_NK)
        r.oz_save_row('zayavki', None, {
            'status': 'work', 'name': subj + TAIL, 'nmck': round(rnd.uniform(2e5, 2e6), 2), 'method': 'ЗЦП',
            'kosgu': rnd.choice(OZ_KOSGU), 'fin': 'Внебюджет', 'resp': rnd.choice(staff), 'initiator': inits,
            'dept': OZ_INIT_DEPT[inits[0]], 'fy': str(Y0)}, None, who())
    # СЗ: отклонённые и в работе без реализации
    for i in range(4):
        inits = rnd.sample(OZ_INIT, 1)
        subj = rnd.choice(SUBJ_NK)
        st = 'отклонена' if i < 2 else 'в работе'
        new_sz(subj, today - timedelta(days=rnd.randint(3, 40)), inits, _items_for(subj, rnd), st=st,
               remarks='Замечания: приложить обоснование закупки у ед. поставщика и 3 коммерческих предложения' if st == 'отклонена' else '')
    # 2) договоры с единственным поставщиком
    ep_nums = []
    for i in range(28):
        cd = date(Y0, 1, 15) + timedelta(days=rnd.randint(0, (today - date(Y0, 1, 15)).days))
        agreed = cd - timedelta(days=rnd.randint(3, 20))
        concluded = i < 23
        if not concluded:
            agreed = today - timedelta(days=rnd.choice([4, 9, 17, 25, 30]))
        at(agreed, 12)
        org = rnd.choice(COMPANIES)[0]
        fin, prog = rnd.choice(OZ_FIN)
        inits = rnd.sample(OZ_INIT, rnd.choice([1, 1, 2]))
        subj = rnd.choice(OZ_EP_SUBJ + ['Поставка компьютерной техники', 'Поставка мебели', 'Поставка лабораторного оборудования'])
        items = _items_for(subj, rnd)
        sz = new_sz(subj, agreed - timedelta(days=rnd.randint(5, 20)), inits, [dict(x) for x in items], st='исполнена' if concluded else 'в работе') if items else None
        s = round(sum(x['sum'] for x in items), 2) if items else round(rnd.uniform(2e4, 6e5), 2)
        cnum = f'{rnd.randint(100, 999)}/ЕП-{cd:%y}'
        if concluded:
            ep_nums.append((cnum, cd, org))
        end_ = cd + timedelta(days=rnd.choice([30, 45, 60]))
        ex = {}
        if concluded:
            if end_ < today:
                late = rnd.random() < .25
                ex = {'ex_delivery': (end_ + timedelta(days=rnd.randint(3, 20)) if late else end_ - timedelta(days=rnd.randint(0, 10))).isoformat(),
                      'ex_status': 'исполнен полностью' if rnd.random() < .7 else 'исполнен частично', 'ex_paid': s if rnd.random() < .7 else round(s * .3, 2),
                      'ex_closed': 'да' if rnd.random() < .5 else 'нет'}
            else:
                ex = {'ex_status': 'исполняется', 'ex_paid': round(s * .3, 2) if rnd.random() < .4 else ''}
        r.oz_save_row('ep', None, dict({
            'sz': sz['sz'] if sz else 'без СЗ', 'sz_links': [sz['id']] if sz else [], 'edo': rnd.choice(['ЭДО', 'ЭДО', 'Бумага']), 'oao_no': f'{rnd.randint(100, 900)}-ЕП',
            'oao_date': agreed.isoformat(), 'kosgu': rnd.choice(OZ_KOSGU), 'org': org, 'cnum': cnum if concluded else '',
            'cdate': cd.isoformat() if concluded else '', 'subject': subj, 'rf': 'да' if any(x['rf'] for x in items) else 'нет', 'sum': s,
            'clause': rnd.choice(['8.1.30', '8.1.30', '8.1.26', '8.1.3']), 'advance': rnd.choice(['без аванса', '30%', '50%', '']),
            'period': rnd.choice(['Одномоментно', 'Партиями']), 'term': f'{(end_ - cd).days} календарных дней', 'validity': date(Y0, 12, 31).isoformat(), 'fy': str(Y0),
            'resp': sz['resp'] if sz else rnd.choice(staff), 'resp_contract': inits, 'dept': OZ_INIT_DEPT[inits[0]],
            'prorector': rnd.choice(OZ_PRORECTORS), 'fin': fin, 'agreement': prog, 'items': [dict(x, src=f"{sz['id']}#{k}") for k, x in enumerate(items)] if sz else []}, **ex), None, who())
    # 3) счета
    for i in range(30):
        dd = date(Y0, 1, 20) + timedelta(days=rnd.randint(0, (today - date(Y0, 1, 20)).days))
        at(dd, 14)
        s = round(rnd.uniform(5e3, 3e5), 2)
        adv = rnd.choice(['100%', '50%', '30%', 'без аванса'])
        paid = s if dd < today - timedelta(days=20) and rnd.random() < .8 else (round(s * int(adv[:-1]) / 100, 2) if adv.endswith('%') and rnd.random() < .5 else '')
        fin, prog = rnd.choice(OZ_FIN)
        inits = rnd.sample(OZ_INIT, 1)
        r.oz_save_row('invoices', None, {
            'sz': f'СЗ {rnd.randint(100, 999)}', 'date_oao': dd.isoformat(), 'kosgu': rnd.choice(OZ_KOSGU), 'org': rnd.choice(COMPANIES)[0],
            'inv': f'№{rnd.randint(1, 400)} от {dd - timedelta(days=rnd.randint(1, 5)):%d.%m.%Y}', 'sum': s, 'subject': rnd.choice(OZ_EP_SUBJ),
            'terms': 'в течение 5 рабочих дней', 'advance': adv, 'paid': paid, 'clause': '8.1.7',
            'paid_date': (dd + timedelta(days=rnd.randint(3, 12))).isoformat() if paid != '' else '', 'resp': rnd.choice(staff),
            'dept_resp': inits, 'dept': OZ_INIT_DEPT[inits[0]], 'fin': fin, 'agreement': prog, 'prorector': rnd.choice(OZ_PRORECTORS), 'fy': str(Y0)}, None, who())
    # 4) доп. соглашения, запросы КП
    for cnum, cd, org in rnd.sample(ep_nums, 6):
        at(cd + timedelta(days=40), 11)
        r.oz_save_row('ds', None, {
            'oao': f'{rnd.randint(100, 900)}-ДС', 'date_out': (cd + timedelta(days=38)).isoformat(), 'cnum': cnum, 'cdate': cd.isoformat(),
            'org': org, 'reason': rnd.choice(['продление срока договора', 'изменение реквизитов', 'изменение цены', 'тех.ошибка']),
            'ds': f'ДС №1 от {cd + timedelta(days=40):%d.%m.%Y}', 'resp': rnd.choice(staff), 'dept_resp': rnd.sample(OZ_INIT, 1)}, None, who())
    for i in range(16):
        dd = date(Y0, 3, 1) + timedelta(days=rnd.randint(0, (today - date(Y0, 3, 1)).days))
        at(dd, 16)
        r.oz_save_row('kp', None, {
            'n': i + 1, 'sent': dd.isoformat(), 'inno': f'вх-{rnd.randint(100, 999)}', 'received': (dd + timedelta(days=rnd.randint(1, 5))).isoformat(),
            'company': rnd.choice(COMPANIES)[0], 'method': rnd.choice(['ЗЦП', 'ЗЭА']), 'price': round(rnd.uniform(1e5, 3e6), 2),
            'subject': rnd.choice(SUBJ_NK + SUBJ_K), 'resp': rnd.choice(staff)}, None, who())
    # 5) нумерация СЗ по порядку поступления
    with r.OZ_LOCK:
        d = r.oz_load('sz')
        for k, row in enumerate(sorted(d['rows'], key=lambda x: x.get('din') or ''), 1):
            row['n'] = k
        r.oz_save_table('sz', d)
    at(today, 9)
    r.oz_recalc('ep')
    print('  ОПАиЗ: сверка с реестром ОАО —', r.oz_sync(force=True), 'строк дополнено')
    # исполнение по заключённым заявкам (в боевой панели придёт из 1С БГУ)
    for z in r.oz_load('zayavki')['rows']:
        if not z.get('cdate') or not z.get('end'):
            continue
        end_ = date.fromisoformat(z['end'])
        s_ = z.get('csum') or 0
        if end_ < today and rnd.random() < .85:
            late = rnd.random() < .2
            ex = {'ex_delivery': (end_ + timedelta(days=rnd.randint(3, 25)) if late else end_ - timedelta(days=rnd.randint(0, 12))).isoformat(),
                  'ex_status': 'исполнен полностью' if rnd.random() < .75 else 'исполнен частично',
                  'ex_paid': s_ if rnd.random() < .75 else round(s_ * .5, 2), 'ex_closed': 'да' if rnd.random() < .6 else 'нет'}
        else:
            ex = {'ex_status': 'исполняется'}
        r.oz_save_row('zayavki', z['id'], ex, None, {'fio': z.get('resp') or 'панель'})
    smp_refs_demo(today)
    opaiz_docs(today)
    # проверка СМП по нескольким договорам ЕП — результат сохраняется в строку
    eps = [x for x in r.oz_load('ep')['rows'] if x.get('items')][:6]
    for x in eps:
        try:
            out = r.smp_check(x['items'], 'договор', x.get('cdate') or today.isoformat())
        except Exception as e:
            print('  проверка СМП:', e)
            break
        items = [dict(it, smp=res['is_smp'], restr=res['restr']) for it, res in zip(x['items'], out['results'])]
        r.oz_save_row('ep', x['id'], {'smpchk': r.smp_summary_text(out['summary']), 'items': items}, None, {'fio': x.get('resp') or 'панель'}, system=True, note='проверка СМП')


def smp_refs_demo(today):
    """Демо-справочники проверки СМП (вымышленные выборки по структуре настоящих файлов)."""
    from openpyxl import Workbook
    d = r.smp_dir()
    d.mkdir(parents=True, exist_ok=True)
    codes = sorted({c for _, items in OZ_CATALOG for _, c, *_ in items} | {'26.20.15.000', '26.20.17.110', '26.20.11.110', '27.20.23.190', '32.50.50.190', '28.29.31.110'})
    wb = Workbook()
    ws = wb.active
    ws.title = 'для работы'
    ws.append(['КОД', 'НАЗВАНИЕ', 'КОД ОТНОСИТСЯ К ПЕРЕЧНЮ (СМП / Б/О)', 'ПРИМЕЧАНИЕ'])
    groups = sorted({c[:5] for c in codes})
    for g in groups:
        ws.append([g, f'Группировка {g} (демо)', 'смп' if g not in ('32.50', '21.20') else 'б/о', ''])
    for c in codes:
        ws.append([c, f'Позиция {c} (демо)', 'б/о' if c.startswith(('32.50.21', '21.20.10')) else 'смп', ''])
    wb.save(d / 'Перечень СМП.xlsx')
    wb = Workbook()
    ws = wb.active
    ws.title = 'Запрет'
    ws.append([])
    ws.append(['Запрет на закупку иностранных товаров (демо-выборка)'])
    ws.append([])
    for i, (nm, c) in enumerate([('Системные блоки (демо)', '26.20.15'), ('Аппараты ИВЛ (демо)', '32.50.21'), ('Мебель для сидения (демо)', '31.01.11.110')], 1):
        ws.append([f'{i}.', nm, c])
    ws2 = wb.create_sheet('Ограничение')
    ws2.append([])
    ws2.append(['Ограничения («второй лишний», демо-выборка)'])
    ws2.append([])
    for i, (nm, c) in enumerate([('Мониторы (демо)', '26.20.17'), ('Перчатки (демо)', '22.19.6'), ('Препараты (демо)', '21.20.10\nза исключением 21.20.10.141'),
                                 ('Оборудование медицинское (демо)', '26.60.12')], 1):
        ws2.append([f'{i}.', nm, c])
    wb.save(d / 'Запрет и ограничение.xlsx')
    wb = Workbook()
    ws = wb.active
    ws.title = str(today.year)
    ws.append(['№ ОАО', 'Дата', 'Организация (Лицо)', 'Предмет договора', 'Сумма', 'Пункт положения', 'Месяц публикации', 'ОКПД 2'])
    m1 = date(today.year, today.month, 1)
    for k, (org, subj, code, cl) in enumerate([('ООО "АЛЬФА"', 'Поставка мониторов', '26.20.17.110', '8.1.30'), ('ООО "БЕТА"', 'Поставка бумаги', '17.12.14.129', '8.1.30'),
                                               ('ООО "ГАММА"', 'Поставка реагентов', '21.20.23.110', '8.1.40 (8.1.30н)'), ('ООО "ДЕЛЬТА"', 'Поставка МФУ', '26.20.18.000', '8.1.30')], 1):
        ws.append([k, m1, org, subj, 100000 * k, cl, m1 if k != 4 else date(today.year - 1, today.month, 1), code])
    wb.save(d / f'Прямые договоры {today.year}.xlsx')



# ---------- ОВРК: полный набор примеров (раньше было отдельное «Демо ОВРК») ----------
OV_ORGS = [('ООО "АЛЬФА"', '0000000001'), ('ООО "БЕТА"', '0000000002'), ('ООО "ГАММА"', '0000000003'), ('ООО "ДЕЛЬТА"', '0000000004'),
        ('ООО "ЭПСИЛОН"', '0000000005'), ('ООО "ДЗЕТА"', '0000000006'), ('ООО "ТЕТА"', '0000000008'), ('ООО "ЙОТА"', '0000000009'),
        ('ООО "КАППА"', '0000000010'), ('ООО "СИГМА"', '0000000011'), ('ИП ПЕТРОВ П.П.', '000000000013'), ('АО "ВОДОКАНАЛ-ДЕМО"', '0000000020'),
        ('АО "ЭНЕРГО-ДЕМО"', '0000000021'), ('ООО "ОМЕГА"', '0000000012')]
OV_SUBJ_K = ['Поставка медицинских изделий (наконечники для дозаторов)', 'Поставка расходных материалов для лаборатории', 'Поставка белья одноразового',
          'Поставка реагентов', 'Поставка лекарственных препаратов', 'Оказание услуг по техническому обслуживанию оборудования',
          'Поставка дезинфицирующих средств', 'Поставка перевязочных материалов']
OV_SUBJ_P = ['Поставка холодного водоснабжения и водоотведения', 'Поставка электрической энергии', 'Оказание услуг по организации питания',
          'Оказание услуг связи', 'Услуги по вывозу отходов', 'Поставка тепловой энергии']



def d2(d):
    return d.strftime('%d.%m.%Y')


def ovrk_full_demo(today):
    STAFF = r.ОВРК_СОТРУДНИКИ
    ORGS, SUBJ_K, SUBJ_P = OV_ORGS, OV_SUBJ_K, OV_SUBJ_P
    pp_dir = r.pp_folder()
    pp_dir.mkdir(parents=True, exist_ok=True)
    rnd = random.Random(1408)
    today = date.today()
    used_nums = set()

    def ppnum():
        while True:
            n = str(rnd.randint(100000, 999999))
            if n not in used_nums:
                used_nums.add(n)
                return n

    def write_pp(num, d, s, inn, payee, purpose, name=None):
        text = ('1CClientBankExchange\nВерсияФормата=1.03\nКодировка=Windows\nОтправитель=Бухгалтерия (демо)\n'
                f'СекцияДокумент=Платежное поручение\nНомер={num}\nДата={d2(d)}\nСумма={s:.2f}\n'
                f'ПлательщикИНН=7704047505\nПлательщик1=Сеченовский Университет (демо)\nПолучательИНН={inn}\nПолучатель1={payee}\n'
                f'НазначениеПлатежа={purpose}\nКонецДокумента\nКонецФайла\n')
        (pp_dir / (name or f'ПП {num} от {d2(d)}.txt')).write_bytes(text.encode('cp1251'))

    contracts = []
    for i in range(24):
        kind = 'прямые' if i >= 16 else 'конкурентные'
        org, inn = ORGS[i % len(ORGS)]
        start = date(Y0, 1, 12) - timedelta(days=rnd.randint(0, 60)) if i < 20 else today - timedelta(days=rnd.randint(20, 60))
        cdate = start - timedelta(days=rnd.randint(20, 50))
        months = rnd.choice([10, 11, 12]) if i < 20 else 6
        term_end = date(Y0, 12, rnd.randint(10, 30)) if kind == 'конкурентные' else date(Y0, 12, 31)
        validity = date(Y0 + 1, 2, 26) if kind == 'конкурентные' else date(Y0, 12, 31)
        price = round(rnd.uniform(4e5, 6e6), 2) if kind == 'конкурентные' else round(rnd.uniform(3e5, 2.5e6), 2)
        cnum = f'{i + 1}{"" if kind == "конкурентные" else "абвгдежз"[i - 16]} от {cdate:%d.%m.%Y}'
        c = {'kind': kind, 'unit': rnd.choice(['КЦ', 'КЦ', 'ВУЗ']), 'subject': rnd.choice(SUBJ_K if kind == 'конкурентные' else SUBJ_P),
             'org': org, 'inn': inn, 'price': price, 'term': f'с {d2(start)} по {d2(term_end)}', 'validity': d2(validity), 'cnum': cnum,
             'monopolist': kind == 'прямые' and 'АО' in org, 'resp': rnd.choice(STAFF), 'upds': []}
        # исполнительные документы: раз в месяц до сегодня
        d, k = start + timedelta(days=rnd.randint(10, 25)), 0
        n_docs = 0 if i == 15 else 99
        while d < today - timedelta(days=2) and k < n_docs:
            k += 1
            s = round(price / months * rnd.uniform(.6, 1.1), 2)
            if sum(u['sum'] for u in c['upds']) + s > price:
                break
            resp = rnd.choice(STAFF)
            pfx = 'акт' if 'услуг' in c['subject'].lower() else 'упд'
            doc = f'{pfx} {rnd.choice(["УТ-", "БП-", ""])}{rnd.randint(100, 9999)} от {d:%d.%m.%y}'
            age = (today - d).days
            u = {'rk': f'{rnd.randint(300, 9900)}-рк' if kind == 'конкурентные' else f'диадок/{d2(d + timedelta(days=3))}', 'resp': resp, 'doc': doc, 'sum': s,
                 'eis': (d + timedelta(days=rnd.randint(2, 8))).isoformat() if age > 14 or rnd.random() < .4 else '', 'pps': []}
            if age > 12:
                pd = d + timedelta(days=rnd.randint(6, min(25, age - 2)))
                parts = [s] if rnd.random() < .85 else [round(s * .6, 2), round(s - round(s * .6, 2), 2)]
                for part in parts:
                    n = ppnum()
                    placed = (today - pd).days > 12 or rnd.random() < .3
                    u['pps'].append({'num': n, 'date': pd.isoformat(), 'sum': part, 'eis': (pd + timedelta(days=rnd.randint(2, 9))).isoformat() if placed else ''})
                    write_pp(n, pd, part, inn, org, f'Оплата по договору № {cnum.split(" от ")[0]} от {cnum.split(" от ")[1]} по {doc}. Без НДС')
            c['upds'].append(u)
            d += timedelta(days=rnd.randint(26, 34))
        contracts.append(c)
    # особые случаи: закрытый, «можно закрыть», истекает срок действия, обязательства просрочены
    c = contracts[3]
    c['status'] = 'закрыт'
    c = contracts[5]   # всё оплачено и внесено — «можно закрыть»
    tot = sum(u['sum'] for u in c['upds'])
    c['price'] = round(tot, 2)
    for u in c['upds']:
        u['eis'] = u['eis'] or (date.fromisoformat(u['pps'][0]['date']) - timedelta(days=3)).isoformat() if u['pps'] else u['eis']
        if not u['pps']:
            n, pd = ppnum(), min(today - timedelta(days=1), date.fromisoformat(r.to_iso(u['doc'])) + timedelta(days=3))
            u['pps'] = [{'num': n, 'date': pd.isoformat(), 'sum': u['sum'], 'eis': (pd + timedelta(days=3)).isoformat()}]
            write_pp(n, pd, u['sum'], c['inn'], c['org'], f'Оплата по договору № {c["cnum"].split(" от ")[0]} по {u["doc"]}')
            u['eis'] = u['eis'] or (pd - timedelta(days=2)).isoformat()
        for p in u['pps']:
            p['eis'] = p['eis'] or (date.fromisoformat(p['date']) + timedelta(days=3)).isoformat()
    contracts[18]['validity'] = d2(today + timedelta(days=18))
    contracts[7]['term'] = f'с 12.01.{Y0} по {d2(today - timedelta(days=12))}'
    # ошибки в платёжках
    def some_pp(ci, ui=0, pi=0):
        return contracts[ci]['upds'][ui]['pps'][pi], contracts[ci]['upds'][ui]
    p, u = some_pp(1, 1)                  # в файле другая сумма
    write_pp(p['num'], date.fromisoformat(p['date']), p['sum'] + 1000, contracts[1]['inn'], contracts[1]['org'], f'Оплата по договору № {contracts[1]["cnum"].split(" от ")[0]} по {u["doc"]}')
    p, u = some_pp(2, 1)                  # в файле другой ИНН получателя
    write_pp(p['num'], date.fromisoformat(p['date']), p['sum'], '0000000099', 'ООО "ПОСТОРОННИЙ"', f'Оплата по договору № {contracts[2]["cnum"].split(" от ")[0]} по {u["doc"]}')
    p, u = some_pp(4, 2)                  # файла нет
    (pp_dir / f'ПП {p["num"]} от {d2(date.fromisoformat(p["date"]))}.txt').unlink(missing_ok=True)
    p, u = some_pp(6, 1)                  # нет номера договора в назначении
    write_pp(p['num'], date.fromisoformat(p['date']), p['sum'], contracts[6]['inn'], contracts[6]['org'], 'Оплата за поставленный товар. Без НДС')
    dup, _ = some_pp(8, 0)                # один номер ПП в двух документах
    contracts[9]['upds'][1]['pps'].append({'num': dup['num'], 'date': dup['date'], 'sum': 1000.0, 'eis': ''})
    p, u = some_pp(10, 1)                 # переплата по документу
    p['sum'] = round(u['sum'] + 2500, 2)
    write_pp(p['num'], date.fromisoformat(p['date']), p['sum'], contracts[10]['inn'], contracts[10]['org'], f'Оплата по договору № {contracts[10]["cnum"].split(" от ")[0]} по {u["doc"]}')
    p, u = some_pp(11, 0)                 # просроченное размещение ПП в ЕИС
    p['eis'] = ''
    p['date'] = (today - timedelta(days=16)).isoformat()
    write_pp(p['num'], today - timedelta(days=16), p['sum'], contracts[11]['inn'], contracts[11]['org'], f'Оплата по договору № {contracts[11]["cnum"].split(" от ")[0]} по {u["doc"]}')
    # платёжки, которые лежат в папке, но ещё не внесены в реестр (панель подскажет, куда их привязать)
    for c in contracts[12:16]:
        u = c['upds'][-1] if c['upds'] else None
        if not u:
            continue
        if u['pps']:
            u['pps'] = []          # платёжки лежат в папке, но в реестр их ещё не внесли
        else:
            n, pd = ppnum(), today - timedelta(days=rnd.randint(1, 3))
            write_pp(n, pd, u['sum'], c['inn'], c['org'], f'Оплата по договору № {c["cnum"].split(" от ")[0]} по {u["doc"]}')
    # имена по-другому — тоже находятся (номер и дата в имени)
    if contracts[13]['upds'] and contracts[13]['upds'][0]['pps']:
        p = contracts[13]['upds'][0]['pps'][0]
        f = pp_dir / f'ПП {p["num"]} от {d2(date.fromisoformat(p["date"]))}.txt'
        if f.exists():
            f.rename(pp_dir / f'пп_{p["num"]}_{date.fromisoformat(p["date"]):%d%m%Y}.txt')
    # доп. соглашение (новая цена и срок) и претензия; договор на следующий год
    contracts[16]['ds'] = [{'num': '№1', 'date': f'{Y0}-05-20', 'reason': 'увеличение объёма', 'price_new': round(contracts[16]['price'] * 1.1, 2),
                            'sum': round(contracts[16]['price'] * 0.1, 2), 'term_new': '', 'validity_new': f'31.01.{Y0 + 1}'}]
    contracts[0]['claims'] = [{'num': '7', 'date': (today - timedelta(days=8)).isoformat(), 'sum': 38_500.0, 'reason': 'просрочка поставки',
                               'status': 'открыта', 'due': (today + timedelta(days=2)).isoformat()}]
    contracts.append({'kind': 'конкурентные', 'unit': 'КЦ', 'subject': 'Поставка медицинских газов на 2027 год', 'org': 'ООО "КАППА"', 'inn': '0000000010',
                      'price': 3_400_000.0, 'term': f'с 01.01.{Y0 + 1} по 31.12.{Y0 + 1}', 'validity': f'31.01.{Y0 + 2}', 'cnum': f'25 от {d2(today - timedelta(days=5))}',
                      'resp': STAFF[0], 'upds': []})
    for c in contracts:
        r.ovrk_save(None, c, None, {'fio': c['resp']}, note='демо: загружен')
    r.pp_index(force=True)


# Специализация (со слов ОПАиЗ, правило нежёсткое) — по листам файла потребности.
SPEC = {'Казначеева Е.Н.': ['оборудование'], 'Небогина О.Ю.': ['реагенты', 'расходка'], 'Григорьева М.Г.': ['инструменты', 'реагенты', 'животные'],
        'Митрофанова М.Э.': ['сувенирка', 'услуги'], 'Недорослева Н.П.': ['услуги']}
SHEET_KW = [('лек. средства', ('натрия хлорид', 'лидокаин', 'препарат', 'раствор для инъекц')), ('реагенты', ('реактив', 'реагент', 'набор для', 'буфер')),
            ('канцелярия', ('бумаг', 'ручк', 'папк', 'канцеляр')), ('животные', ('мыш', 'крыс', 'кролик', 'животн')),
            ('оборудование', ('центрифуг', 'шкаф ламинар', 'весы', 'термостат', 'тренажёр', 'панель', 'мфу', 'ноутбук', 'монитор', 'принтер', 'системный блок', 'сканер', 'оборудован', 'анализатор')),
            ('расходка', ('пробирк', 'наконечник', 'перчатк', 'шприц', 'бинт', 'салфетк', 'посуд', 'маск'))]


def _cat(name):
    n = ' ' + str(name).lower() + ' '
    return next((c for c, kws in SHEET_KW if any(k in n for k in kws)), 'расходка')


def demo_samples(today):
    """Образцы для показа: таблица ОАО с маршрутом согласования и перечень платёжек от бухгалтерии; специализация ОПАиЗ."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    rnd = random.Random(2110)
    sysw = {'fio': 'Сальникова Е.А.'}
    # специализация + распределение нескольких СЗ по листам (как делает начальник)
    lists = r.oz_lists()
    lists['spec'] = [f"{k}: {', '.join(v)}" for k, v in SPEC.items()]
    r.oz_save_lists(lists)
    load, n = {}, 0
    for row in sorted(r.oz_load('sz')['rows'], key=lambda x: x.get('din') or '', reverse=True):
        items = row.get('items') or []
        if row.get('st') != 'в работе' or not items:
            continue
        for it in items:
            it['cat'] = _cat(it.get('name'))
        if len({it['cat'] for it in items}) < 2 and n >= 3:
            continue
        for it in items:
            who = sorted([f for f, cs in SPEC.items() if it['cat'] in cs], key=lambda f: load.get(f, 0))
            it['resp'] = who[0] if who else 'Сальникова Е.А.'
            load[it['resp']] = load.get(it['resp'], 0) + 1
        r.oz_save_row('sz', row['id'], {'items': items}, None, sysw, note='распределено по листам')
        n += 1
        if n >= 8:
            break
    # 1С ДО: статус согласования у заявок, ещё не ушедших в ОАО
    stages = ['руководитель подразделения', 'ПЭО', 'ОАО', 'юристы', 'бухгалтерия', 'проректор']
    who_by = {'руководитель подразделения': 'Иванов И.И.', 'ПЭО': 'Петрова С.В.', 'ОАО': 'Гончарова К.А.', 'юристы': 'Смирнова А.А.', 'бухгалтерия': 'Кузьмина Е.В.', 'проректор': 'Рязанцев И.И.'}
    for k, z in enumerate([z for z in r.oz_load('zayavki')['rows'] if not z.get('pub') and z.get('status') in ('sign', 'work')]):
        st = stages[k % len(stages)]
        r.oz_save_row('zayavki', z['id'], {'do_stage': st, 'do_who': who_by[st], 'do_since': r.back_workdays(today, [2, 8, 3, 12, 1, 4][k % 6]).isoformat() if hasattr(r, 'back_workdays') else (today - timedelta(days=[2, 10, 4, 16, 1, 5][k % 6])).isoformat()}, None, sysw)
    # новости
    news = [('ОАО', True, 'С 13.10 извещения принимаем только в .docx', 'Сканы и .pdf больше не берём: панель читает Word сама. Если извещение пришло в PDF — попросите у площадки Word или вставьте текст в форму.', 2),
            ('ОАО', False, 'Протоколы итогов — в папку «Протоколы\\входящие» своего года', 'Панель вносит их сама; спорные — во вкладке «На проверке».', 5),
            ('ОПАиЗ', False, 'Проверка СМП — до отправки комплекта в 1С ДО', 'Проверяем коды ОКПД2 в панели (ОПАиЗ → Проверка СМП) и сохраняем результат в строку заявки.', 9)]
    for dept, imp, title, text, ago in news:
        it = r.news_save({'title': title, 'text': text, 'dept': dept, 'important': imp, 'since': (today + timedelta(days=2)).isoformat() if imp else ''}, {'fio': 'Гончарова К.А.' if dept == 'ОАО' else 'Сальникова Е.А.'})
        d = r.news_load()
        for x in d['items']:
            if x['id'] == it['id']:
                x['at'] = (datetime.combine(today - timedelta(days=ago), datetime.min.time()) + timedelta(hours=9, minutes=15)).isoformat(timespec='seconds')
        r.write_atomic(r.ПАПКА_ПАНЕЛИ / 'Новости.json', r.json.dumps(d, ensure_ascii=False, indent=1))
    hfont, hfill = Font(bold=True), PatternFill('solid', fgColor='DDE6F0')

    def sheet(ws, head, rows, widths):
        ws.append(head)
        for c in ws[1]:
            c.font, c.fill, c.alignment = hfont, hfill, Alignment(wrap_text=True, vertical='top')
        for r_ in rows:
            ws.append(r_)
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[ws.cell(1, i).column_letter].width = w
    d = lambda x: (today - timedelta(days=x)).strftime('%d.%m.%Y')
    # --- 1. таблица ОАО: маршрут согласования ---
    Z = r.oz_load('zayavki')['rows']
    pend = [z for z in Z if not z.get('pub') and z.get('status') in ('sign', 'work')]
    pubd = [z for z in Z if z.get('pub')][-3:]
    rows = []
    for k, z in enumerate(pend):
        a = 4 + k * 3
        st = [d(a + 14), f'{d(a + 13)}-{d(a + 10)}', f'{d(a + 10)}-{d(a + 7)}' if k % 3 else d(a + 9), d(a + 2) if k % 3 else '', '', '', '']
        if k == 1:
            st = [d(26), f'{d(25)}-{d(21)}', f'{d(21)}-{d(20)}', f'{d(20)}-{d(9)} замеч. по проекту договора', d(9), '', '']   # долго в бухгалтерии
        if k == 2:
            st = [d(18), f'{d(17)}-{d(15)}', f'{d(15)}-{d(14)}', f'{d(14)}-{d(12)}', f'{d(12)}-{d(10)}', '', d(9)]           # вернулась, ждёт публикации
        rows.append([rnd.randint(1500, 2600), st[0], z.get('n1c'), z.get('name'), z.get('nmck'), (z.get('resp') or '') + ' 50-' + str(rnd.randint(10, 99)),
                     z.get('method') or 'ЗЦП', st[1], st[2], st[3], st[4], st[5], st[6], ''])
    for z in pubd:   # уже опубликованные — панель покажет «опубликована», менять нечего
        rows.append([rnd.randint(1200, 1500), d(40), z.get('n1c'), z.get('name'), z.get('nmck'), z.get('resp') or '', z.get('method') or '', f'{d(39)}-{d(36)}', f'{d(36)}-{d(35)}',
                     f'{d(35)}-{d(33)}', f'{d(33)}-{d(31)}', '', d(30), ''])
    rows.append([2701, d(6), '9999', 'Поставка учебных пособий (лот № 3) — заявки нет в реестре ОПАиЗ', 380000, 'Иванова И.И. 50-11', 'ЗЦП', f'{d(5)}-{d(3)}', d(3), '', '', '', '', ''])
    if pend:
        z = pend[-1]
        rows.append([2702, d(8), '', (z.get('name') or '').replace('Поставка', 'Закупка на поставку', 1), z.get('nmck'), z.get('resp') or '', 'ЗЦП', f'{d(7)}-{d(6)}', d(6), '', '', '', '', 'без номера 1С — панель подберёт по названию'])
    head = ['№ ОАО', 'Дата', '№ заявки 1С', 'Предмет закупки', 'НМЦК', 'Исполнитель в службах/тел', 'Способ', 'ПЭО передача/возврат',
            'Дата передачи ответственному в ОАО (223-ФЗ, СМП)', 'Юр. отдел передача/возврат', 'Бухгалтерия передача/возврат', 'Проректору передача/возврат',
            'На публикацию ответственному', 'Замечания/примечания']
    wb = Workbook()
    sheet(wb.active, head, rows, [8, 11, 11, 60, 14, 22, 9, 22, 26, 26, 18, 16, 30])
    wb.active.title = 'Процедуры'
    out = DOCS / '6 ОПАиЗ — маршрут согласования (таблица ОАО)'
    shutil.rmtree(out, ignore_errors=True)
    out.mkdir(parents=True, exist_ok=True)
    wb.save(out / 'Таблица ОАО — процедуры (маршрут согласования).xlsx')
    # --- 2. перечень платёжек от бухгалтерии ---
    data, _, _ = r.ovrk_all()
    pp_dir = r.pp_folder()
    pays, used = [], set()

    def ppn():
        while True:
            x = str(rnd.randint(100000, 999999))
            if x not in used:
                used.add(x)
                return x
    open_c = [c for c in data['rows'] if c.get('status') != 'закрыт' and c.get('upds') and c.get('inn')]
    for c in open_c:
        u = next((x for x in reversed(c['upds']) if isinstance(x.get('rest'), (int, float)) and x['rest'] > 1), None)
        if not u:
            continue
        if sum(1 for x in pays if x[5].startswith('Оплата по договору')) >= 5:
            break
        cn = str(c['cnum']).split(' от ')
        pays.append([ppn(), today - timedelta(days=rnd.randint(1, 4)), round(u['rest'], 2), c.get('org'), c.get('inn'),
                     f"Оплата по договору № {cn[0]} от {cn[1] if len(cn) > 1 else ''} по {u.get('doc') or ''}. Без НДС"])
    if open_c:   # только ИНН, без номера договора — аванс
        c = open_c[-1]
        pays.append([ppn(), today - timedelta(days=2), 150000.0, c.get('org'), c.get('inn'), 'Авансовый платёж за поставку товара. Без НДС'])
    dup = next(((c, p) for c in data['rows'] for u in c.get('upds') or [] for p in u.get('pps') or [] if p.get('num') and p.get('date')), None)
    if dup:
        c, p = dup
        pays.append([p['num'], date.fromisoformat(p['date']), p.get('sum') or 0, c.get('org'), c.get('inn'), f"Оплата по договору № {str(c['cnum']).split(' от ')[0]}"])
    pays.append([ppn(), today - timedelta(days=1), 48200.0, 'ООО "МАЛЫЙ ДОГОВОР"', '0000000077', 'Оплата по счёту № 15 от 01.10.2026 за канцтовары'])
    for num, dt, sm, org, inn, purp in pays:
        if not (pp_dir / f'ПП {num} от {dt:%d.%m.%Y}.txt').exists():
            txt = ('1CClientBankExchange\nВерсияФормата=1.03\nКодировка=Windows\nСекцияДокумент=Платежное поручение\n'
                   f'Номер={num}\nДата={dt:%d.%m.%Y}\nСумма={sm:.2f}\nПлательщикИНН=7704047505\nПолучательИНН={inn}\nПолучатель1={org}\n'
                   f'НазначениеПлатежа={purp}\nКонецДокумента\nКонецФайла\n')
            (pp_dir / f'ПП {num} от {dt:%d.%m.%Y}.txt').write_bytes(txt.encode('cp1251'))
    wb = Workbook()
    sheet(wb.active, ['№ п/п', 'Номер ПП', 'Дата ПП', 'Сумма', 'Получатель', 'ИНН получателя', 'Назначение платежа'],
          [[i + 1, x[0], x[1], x[2], x[3], x[4], x[5]] for i, x in enumerate(pays)], [6, 11, 11, 14, 30, 14, 70])
    wb.active.title = 'Платёжки'
    for row in wb.active.iter_rows(min_row=2):
        row[2].number_format = 'DD.MM.YYYY'
        row[3].number_format = '#,##0.00'
    out = DOCS / '5 ОВРК — платёжки бухгалтерии и таблицы ОАО за неделю'
    shutil.rmtree(out, ignore_errors=True)
    out.mkdir(parents=True, exist_ok=True)
    wb.save(out / f'Перечень платёжек (бухгалтерия) {today:%d.%m.%Y}.xlsx')
    # --- 3. файл потребности к СЗ (по листам) ---
    need = {'расходка': [('Наконечники для дозаторов 200 мкл', 'уп', 20, 650), ('Пробирки микроцентрифужные 1,5 мл', 'уп', 10, 540), ('Перчатки нитриловые, размер M', 'пара', 1000, 9.5)],
            'оборудование': [('Центрифуга лабораторная настольная', 'шт', 1, 185000), ('Термостат суховоздушный', 'шт', 1, 98000)],
            'реагенты': [('Набор реагентов для выделения ДНК', 'набор', 4, 32000), ('Буферный раствор PBS', 'фл', 10, 1800)],
            'лек. средства': [('Натрия хлорид р-р 0,9% 500 мл', 'фл', 200, 65)],
            'животные': [('Мыши линии BALB/c, самки 6–8 нед.', 'гол', 40, 900)],
            'канцелярия': [('Бумага офисная А4', 'пач', 50, 420), ('Папка-регистратор', 'шт', 30, 210)]}
    wb = Workbook()
    wb.remove(wb.active)
    for sh, rows_ in need.items():
        ws = wb.create_sheet(sh)
        ws.append([f'Потребность к служебной записке · лист «{sh}»'])
        sheet_rows = [[k + 1, nm, un, q, pr, round(q * pr, 2)] for k, (nm, un, q, pr) in enumerate(rows_)]
        ws.append([])
        for c_, h in enumerate(['№', 'Наименование', 'Ед. изм.', 'Кол-во', 'Цена за ед., руб.', 'Сумма, руб.'], 1):
            ws.cell(3, c_, h).font = hfont
        for rr in sheet_rows:
            ws.append(rr)
        for c_, w in zip('ABCDEF', (5, 48, 10, 10, 16, 16)):
            ws.column_dimensions[c_].width = w
    wb.save(DOCS / '4 ОПАиЗ — перетащите на вкладку ОПАиЗ' / 'Потребность к СЗ (по листам).xlsx')
    # --- 4. таблицы ОАО за неделю для ОВРК: новые заключённые договоры и ДС ---
    data, _, _ = r.ovrk_all()
    old = [c for c in data['rows'] if c.get('cnum') and ' от ' in c['cnum']][:3]
    wk = lambda k: (today - timedelta(days=k)).strftime('%d.%m.%Y')
    head = ['№ ОАО', 'Дата', '№ КЦ/д-26', 'Организация (Лицо)', 'Предмет договора', 'Сумма', 'ИНН', 'Срок поставки/оказания услуг/выполнения работ', 'Срок действия договора', 'Номер договора']
    proc = [[2611, wk(20), '', 'ООО "ГАММА"', 'Поставка лабораторного оборудования (лот № 40)', 2_480_000.0, '0000000003', f'с {wk(3)} по 31.12.{today.year}', f'31.01.{today.year + 1}', f'32616455001 от {wk(3)}'],
            [2615, wk(18), '1/КЦ', 'ООО "ДЕЛЬТА"', 'Поставка реагентов (лот № 41)', 1_150_400.0, '0000000004', f'с {wk(2)} по 30.11.{today.year}', f'31.12.{today.year}', f'32616455002 от {wk(2)}'],
            [2620, wk(15), '', 'ООО "СИГМА"', 'Поставка расходных материалов (лот № 42)', 312_500.0, '0000000011', f'с {wk(1)} по 30.11.{today.year}', f'31.12.{today.year}', f'COM0410260000{rnd.randint(10, 99)} от {wk(1)}'],
            [2622, wk(14), '', 'ООО "ЭПСИЛОН"', 'Выполнение работ по ремонту помещений (лот № 43)', 4_870_000.0, '0000000005', f'с {wk(4)} по 31.12.{today.year}', f'31.03.{today.year + 1}', f'32616455003 от {wk(4)}']]
    if old:
        c = old[0]
        proc.append([1500, wk(200), '', c.get('org'), c.get('subject'), c.get('price'), c.get('inn'), c.get('term'), c.get('validity'), c.get('cnum')])
    direct = [[3101, wk(10), '', 'АО "ЭНЕРГО-ДЕМО"', 'Поставка электрической энергии (доп. корпус)', 1_920_000.0, '0000000021', f'с {wk(5)} по 31.12.{today.year}', f'31.12.{today.year}', f'ЭС-77/26 от {wk(5)}'],
              [3102, wk(9), '', 'ИП ПЕТРОВ П.П.', 'Услуги по вывозу отходов', 640_000.0, '000000000013', f'с {wk(2)} по 31.12.{today.year}', f'31.12.{today.year}', f'ВО-12 от {wk(2)}']]
    ds = [['№ ДС', 'Дата поступления', 'Ответственный в службах', 'ВУЗ/КЦ', '№ договора в ОАО', '№ и дата заключения договора', '№ ДС (ДС №1, ДС № 2)', 'Организация (лицо)', 'Причина заключения ДС', 'Сумма ДС', 'Дата заключения ДС']]
    for k, c in enumerate(old[:2]):
        ds.append([k + 1, wk(6), 'Иванова И.И.', 'ВУЗ', 100 + k, c.get('cnum'), f'№{len(c.get("ds") or []) + 1}', c.get('org'), 'увеличение объёма' if k == 0 else 'продление срока', 150_000.0 if k == 0 else '', wk(3)])
    wb = Workbook()
    sheet(wb.active, head, proc, [8, 11, 9, 26, 46, 14, 13, 34, 16, 26])
    wb.active.title = 'Процедуры'
    sheet(wb.create_sheet('Прямые'), head, direct, [8, 11, 9, 26, 46, 14, 13, 34, 16, 26])
    sheet(wb.create_sheet('ДС'), ds[0], ds[1:], [6, 12, 18, 8, 10, 26, 12, 26, 20, 12, 14])
    out = DOCS / '5 ОВРК — платёжки бухгалтерии и таблицы ОАО за неделю'
    wb.save(out / f'Таблицы ОАО за неделю — заключённые договоры и ДС {today:%d.%m.%Y}.xlsx')
    r.pp_index(force=True)


def demo_extras(today):
    """Доработки 10.10: каталог для инициаторов, годы, «всё ок?», РФ-спецификация, ОВРК со связью с ОПАиЗ, журнал уведомлений."""
    import re as _re
    rnd = random.Random(1010)
    sysw = {'fio': 'Недорослева Н.П.'}
    # 1) каталог позиций (результат spec_extract.py по обезличенным спецификациям)
    src = HERE / 'Каталог (демо)'
    r.ПАПКА_КАТАЛОГА.mkdir(parents=True, exist_ok=True)
    for f in ('Каталог.html', 'Каталог позиций.xlsx'):
        if (src / f).exists():
            shutil.copy(src / f, r.ПАПКА_КАТАЛОГА / f)
    # 2) ОПАиЗ
    r.oz_sync(force=True)
    sz = sorted(r.oz_load('sz')['rows'], key=lambda x: x.get('din') or '')
    for row in sz[-4:]:                               # служебные записки уже на 2027 год
        r.oz_save_row('sz', row['id'], {'fy': str(Y0 + 1), 'subject': row.get('subject', '') + ' (на 2027 год)'}, None, sysw)
    for row in rnd.sample(sz, min(14, len(sz))):
        r.oz_save_row('sz', row['id'], {'complexity': rnd.choice(['низкая', 'средняя', 'средняя', 'высокая']), 'urgency': rnd.choice(['низкая', 'средняя', 'высокая'])}, None, sysw)
    z = sorted(r.oz_load('zayavki')['rows'], key=lambda x: x.get('d1c') or '')
    for row in z[-2:]:
        r.oz_save_row('zayavki', row['id'], {'fy': str(Y0 + 1)}, None, sysw)
    inv = r.oz_load('invoices')['rows']
    old = [x for x in inv if isinstance(x.get('oklong'), (int, float)) and x['oklong'] >= r.СЧЁТ_ПРОВЕРКА_РАБ_ДНЕЙ and isinstance(x.get('rest'), (int, float)) and x['rest'] > 0.01]
    for k, row in enumerate(old):
        if k < 2:
            continue                                  # два счёта — без ответа: вопрос «всё ок?» в «Моём дне»
        r.oz_save_row('invoices', row['id'], {'okcheck': 'не ок' if k == 2 else 'ок'}, None, sysw)
    for row in inv:
        t = str(row.get('terms') or '').lower()
        pt = 'аванс — 5 рабочих дней' if 'аванс' in t or 'предоплат' in t else 'постоплата — 7 рабочих дней'
        r.oz_save_row('invoices', row['id'], {'pay_terms': pt, 'terms': _re.sub(r'(?i).*?(срок[^;]*)$', r'\1', row.get('terms') or '') if ';' in str(row.get('terms') or '') else ''}, None, sysw)
    for row in rnd.sample(inv, min(4, len(inv))):
        r.oz_save_row('invoices', row['id'], {'kosgu': [row.get('kosgu') or '346-расх.материалы', '226-услуги/работы']}, None, sysw)
    for t in ('ep', 'zayavki'):                      # РФ-оборудование: договор есть, спецификация не загружена → задача
        cand = [x for x in r.oz_load(t)['rows'] if x.get('cnum') and x.get('cdate')]
        if cand:
            r.oz_save_row(t, cand[-1]['id'], {'rf': 'да', 'items': []}, None, sysw, system=True)
    ovrk_full_demo(today)
    # 3) ОВРК: договоры > 500 000 из ОПАиЗ (исполнение = сумма ПП) + прямые с ДС и претензией
    pp_dir = r.pp_folder()
    pp_dir.mkdir(parents=True, exist_ok=True)
    used = set()

    def ppnum():
        while True:
            n = str(rnd.randint(100000, 999999))
            if n not in used:
                used.add(n)
                return n

    def write_pp(num, d, s_, inn, payee, purpose):
        text = ('1CClientBankExchange\nВерсияФормата=1.03\nКодировка=Windows\nСекцияДокумент=Платежное поручение\n'
                f'Номер={num}\nДата={d:%d.%m.%Y}\nСумма={s_:.2f}\nПлательщикИНН=7704047505\nПолучательИНН={inn}\nПолучатель1={payee}\n'
                f'НазначениеПлатежа={purpose}\nКонецДокумента\nКонецФайла\n')
        (pp_dir / f'ПП {num} от {d:%d.%m.%Y}.txt').write_bytes(text.encode('cp1251'))

    staff = r.ОВРК_СОТРУДНИКИ
    big = []
    for t, amt_k in (('zayavki', 'csum'), ('ep', 'sum')):
        for x in r.oz_load(t)['rows']:
            amt = x.get(amt_k)
            if isinstance(amt, (int, float)) and amt > r.ОВРК_ПОРОГ_ИСПОЛНЕНИЯ and x.get('cnum') and r.re.fullmatch(r'\d{4}-\d\d-\d\d', str(x.get('cdate') or '')):
                big.append((t, x, amt))
    for k, (t, x, amt) in enumerate(big[:14]):
        cd = date.fromisoformat(x['cdate'])
        org = str(x.get('supplier') or x.get('org') or 'ООО "АЛЬФА"')
        inn = next((c[1] for c in COMPANIES if c[0] == org), '')
        cnum = f"{x['cnum']} от {cd:%d.%m.%Y}"
        c = {'kind': 'неконкурентные' if t == 'zayavki' and k % 3 == 0 else ('прямые' if t == 'ep' else 'конкурентные'), 'unit': rnd.choice(['КЦ', 'ВУЗ']),
             'subject': str(x.get('name') or x.get('subject') or '')[:200], 'org': org, 'inn': inn, 'price': round(amt, 2),
             'term': f'с {cd:%d.%m.%Y} по 31.12.{Y0}', 'validity': f'31.12.{Y0}', 'cnum': cnum, 'upds': []}
        share = [0, .35, .6, 1.0][k % 4]                # исполнено: 0 %, ~35 %, ~60 %, 100 %
        n_upd = 0 if not share else (1 if share == 1.0 else 2)
        d = cd + timedelta(days=12)
        for j in range(n_upd):
            if d >= today:
                break
            s_ = round(amt * share / n_upd, 2)
            u = {'rk': f'{rnd.randint(7000, 9000)}-рк', 'resp': rnd.choice(staff), 'doc': f'упд {rnd.randint(100, 9999)} от {d:%d.%m.%y}', 'sum': s_,
                 'eis': (d + timedelta(days=3)).isoformat(), 'pps': []}
            pd = min(today - timedelta(days=1), d + timedelta(days=8))
            num = ppnum()
            u['pps'].append({'num': num, 'date': pd.isoformat(), 'sum': s_, 'eis': (pd + timedelta(days=3)).isoformat() if (today - pd).days > 6 else ''})
            write_pp(num, pd, s_, inn, org, f"Оплата по договору № {x['cnum']} от {cd:%d.%m.%Y}")
            c['upds'].append(u)
            d += timedelta(days=30)
        r.ovrk_save(None, c, None, {'fio': rnd.choice(staff)}, note='демо: договор > 500 000 из ОПАиЗ')
    r.pp_index(force=True)
    demo_samples(today)
    # 4) уведомления о публикациях за последние 10 дней (извещения записались при внесении)
    r.cache['mtime'] = None
    r.refresh_cache(force=True)
    lim = (today - timedelta(days=10)).isoformat()
    for x in sorted([x for x in r.cache['rows'] if x.get('pub') and x['pub'] >= lim and x.get('number')], key=lambda x: x['pub']):
        when = datetime.fromisoformat(x['pub'] + 'T10:%02d:00' % rnd.randint(0, 59))
        r.notify('pub', {'oao': x.get('oao'), 'number': x.get('number'), 'name': x.get('name'), 'nmck': x.get('nmck'), 'pub': x.get('pub'),
                         'bid': x.get('bid'), 'responsible': x.get('responsible')}, when=when)
    data = r.notify_load()
    data['pub_seen'] = sorted({x.get('number') for x in data.get('items', []) if x.get('kind') == 'pub'})
    data['items'] = sorted(data.get('items', []), key=lambda x: x.get('ts', ''))
    r.write_atomic(r._notify_path(), r.json.dumps(data, ensure_ascii=False, indent=1))


def opaiz_docs(today):
    """Документы для показа реестра ОПАиЗ и проверки СМП."""
    from openpyxl import Workbook
    from docx import Document
    out = DOCS / '4 ОПАиЗ — перетащите на вкладку ОПАиЗ'
    shutil.rmtree(out, ignore_errors=True)
    out.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = 'Заявка'
    for line in [['Утверждаю'], ['Первый проректор'], [''], ['Заявка на закупку неконкурентным способом запрос ценовых предложений в электронной форме'],
                 [f'Дата составления: {today:%d.%m.%Y}'], ['План-график: п. № 41'], ['Заявка 1С', '2528'], ['Код по ЭК 310'], ['Площадка РосЭлТорг'],
                 ['Предмет закупки', 'Сумма, руб.'],
                 ['Закупка неконкурентным способом запрос ценовых предложений в электронной форме на поставку лабораторного оборудования' + TAIL + ' с субъектом малого и среднего предпринимательства.', 1_250_000.00],
                 ['№', 'Наименование, краткие характеристики', 'Технические характеристики', 'Ед. изм.', 'Кол-во', 'код ОКПД 2', 'Постановление 1875', 'Цена за ед. изм., руб.', 'Ставка НДС, %', 'Сумма с учетом НДС, руб.', 'Источник финансирования', '', 'Подразделение и ответственные'],
                 ['', '', '', '', '', '', '', '', '', '', 'софинансирование', 'грант'],
                 [1, 'Центрифуга лабораторная настольная', 'условные характеристики', 'шт', 2, '28.99.39.190', 'Перечень 2, применяется преимущество', 185000, 'НДС не облагается', 370000, 'Грант', 'Приоритет', '01.000.Е.129 Орлова НЦ'],
                 [2, 'Шкаф ламинарный II класса', 'условные характеристики', 'шт', 1, '28.25.14.110', 'не применяется', 640000, 'НДС не облагается', 640000, '', 'Приоритет', ''],
                 [3, 'Весы аналитические', 'условные характеристики', 'шт', 1, '28.29.31.110', 'не применяется', 240000, 'НДС не облагается', 240000, '', 'Приоритет', ''],
                 ['ИТОГО', 1_250_000],
                 ['Место доставки: г. Москва (склад). Ответственные лица: Директор Научного центра Орлова Елена Николаевна, телефон 000; Контактное лицо - Заведующий лабораторией № 1 Иванов Иван Иванович'],
                 ['Условия поставки: Поставка Товара производится Заказчику партиями, в течение 30 рабочих дней с момента заключения Договора'],
                 [f'Срок действия договора: с момента заключения договора по 31.12.{today.year}'],
                 ['Преимущества, предоставляемые заказчиком СМП - Требуются'],
                 ['Исполнитель: Казначеева Е.Н.']]:
        ws.append(line)
    wb.save(out / 'Заявка на закупку 2528 (ЗЦП, новый шаблон).xlsx')
    doc = Document()
    for t in ['ОБОСНОВАНИЕ', 'цены договора с единственным поставщиком', 'Характеристики объекта закупки',
              'Поверка средств измерений лаборатории (условный перечень) для нужд университета',
              'Обоснование способа закупки', 'Закупка у единственного поставщика по п. 8.1 Положения о закупке', 'Доход научного центра',
              'Расчет цены контракта', f'ООО "ДЕЛЬТА" Счет №58 от {today - timedelta(days=3):%d.%m.%Y} на сумму 64 800, 00 руб.',
              'Сроки оплаты', 'Оплата счёта по 100% предоплата в течение 5 рабочих дней', 'Срок оказания услуг: 20 календарных дней',
              f'Дата подготовки обоснования: {today:%d.%m.%Y}', 'Ответственный: Орлова Е.Н.', 'Исполнитель: Небогина О.Ю.']:
        doc.add_paragraph(t)
    # как в отделе: таблица «вопрос | ответ» с ОКПД2 — по ней работает и проверка СМП
    tb = doc.add_table(rows=0, cols=2)
    tb.style = 'Table Grid'
    for k, v in [('Предмет договора', 'Поверка средств измерений лаборатории (условный перечень)'), ('Код ОКПД 2', '71.20.19.130; 26.51.66.190'),
                 ('Сроки оплаты', '100% предоплата в течение 5 рабочих дней'), ('Сроки поставки/установки/монтажа / Сроки оказания услуг, выполнения работ', '20 календарных дней'),
                 ('Расчет цены контракта', f'ООО "ДЕЛЬТА" Счет №58 от {today - timedelta(days=3):%d.%m.%Y} на сумму 64 800,00 руб.')]:
        cs = tb.add_row().cells
        cs[0].text, cs[1].text = k, v
    doc.save(out / 'Обоснование к счёту — ДЕЛЬТА.docx')

    def contract(fname, subj_head, obj, rows, total, clause='8.1.30'):
        doc = Document()
        for t in ['Д О Г О В О Р', subj_head + ' №', 'г. Москва «____» ___________ 20__ года',
                  'Федеральное государственное автономное образовательное учреждение высшего образования Первый Московский государственный медицинский '
                  'университет имени И.М. Сеченова Министерства здравоохранения Российской Федерации (Сеченовский Университет), именуемое в дальнейшем «Заказчик», '
                  'в лице Директора Клинического центра Волковой Ольги Сергеевны, действующей на основании Доверенности, с одной стороны, и '
                  'Общество с ограниченной ответственностью «ДЕЛЬТА», именуемое в дальнейшем «Поставщик», в лице Генерального директора, '
                  'действующего на основании Устава, с другой стороны, заключили настоящий Договор о нижеследующем:',
                  f'1.1. Поставщик обязуется передать Заказчику на условиях настоящего Договора {obj} (далее – Товар) по наименованиям, в количестве и с характеристиками согласно Спецификации (Приложение № 1).',
                  f'1.2. Закупка производится на основании пункта {clause} Положения о закупке товаров, работ, услуг.',
                  f'3.1. Цена Договора составляет {total:,.0f} (условная сумма) рублей 00 копеек, НДС не облагается.'.replace(',', ' '),
                  '4.1. Срок поставки: в течение 20 рабочих дней с даты заключения Договора.',
                  'Поставка осуществляется партиями по заявкам Заказчика.',
                  f'9.1. Договор вступает в силу с момента его заключения и действует по 31.12.{today.year}',
                  'Договор заключён в рамках программы «Приоритет-2030», ответственный по договору Белов С.А.',
                  'Поставщик: ООО «ДЕЛЬТА» ИНН 0000000004', '__________ Волкова О.С.', 'Приложение № 1. Спецификация']:
            doc.add_paragraph(t)
        tb = doc.add_table(rows=1, cols=7)
        for i, h in enumerate(['№', 'Наименование продукции', 'Количество', 'Ед. изм.', 'Цена, руб.', 'Сумма, руб.', 'код ОКПД2']):
            tb.rows[0].cells[i].text = h
        for k, (nm, q, u, p, code) in enumerate(rows, 1):
            c = tb.add_row().cells
            for i, v in enumerate([k, nm, q, u, p, q * p, code]):
                c[i].text = str(v)
        doc.save(out / fname)

    contract('Договор ЕП — ДЕЛЬТА (реактивы).docx', 'на поставку реактивов', 'реактивы', [('Набор реагентов ИФА', 5, 'набор', 24000, '21.20.23.110'), ('Буферный раствор', 18, 'фл', 3800, '20.59.52.199'),
             ('Натрия хлорид р-р 0,9%', 400, 'фл', 54, '21.20.10.141')], 5 * 24000 + 18 * 3800 + 400 * 54)
    contract('Договор ЕП — ДЕЛЬТА (компьютерная техника, для проверки СМП).docx', 'поставки', 'компьютерную технику и мебель',
             [('Системный блок', 4, 'шт', 78000, '26.20.15.000'), ('Монитор 27"', 4, 'шт', 26000, '26.20.17.110'), ('МФУ лазерное', 1, 'шт', 48000, '26.20.18.000'),
              ('Стол лабораторный', 2, 'шт', 46000, '31.01.12.150'), ('Кресло офисное', 4, 'шт', 9800, '31.01.11.110'), ('Перчатки нитриловые', 500, 'пар', 18, '22.19.60.119')],
             4 * 78000 + 4 * 26000 + 48000 + 2 * 46000 + 4 * 9800 + 500 * 18)


def demo_smp_categories():
    """В демо у договоров проставлена категория победителя («Кат. см и сп») — для аналитики по СМП."""
    from openpyxl import load_workbook as _lw
    rnd = random.Random(55)
    for y in (Y0 - 1, Y0):
        p = REG(y)
        wb = _lw(p)
        for ws in wb.worksheets:
            hdr = None
            for i in range(1, 6):
                vals = {str(c.value or '').strip(): c.column for c in ws[i]}
                if 'Кат. см и сп' in vals:
                    hdr = (i, vals['Кат. см и сп'], next((c for k, c in vals.items() if k.startswith('Номер и дата договора')), None))
                    break
            if not hdr or not hdr[2]:
                continue
            for row in range(hdr[0] + 1, ws.max_row + 1):
                if ws.cell(row, hdr[2]).value and not ws.cell(row, hdr[1]).value:
                    ws.cell(row, hdr[1]).value = rnd.choices(['микро', 'малое', 'среднее', 'СОНКО', 'не СМП'], [30, 34, 10, 3, 23])[0]
        wb.save(p)

MONTHS_RU = ['Январь', 'Февраль', 'Март', 'Апрель', 'Май', 'Июнь', 'Июль', 'Август', 'Сентябрь', 'Октябрь', 'Ноябрь', 'Декабрь']


def demo_users():
    import json
    r.ПАПКА_ПАНЕЛИ.mkdir(parents=True, exist_ok=True)
    (r.ПАПКА_ПАНЕЛИ / r.ФАЙЛ_ПОЛЬЗОВАТЕЛЕЙ).write_text(json.dumps(DEMO_USERS, ensure_ascii=False, indent=1), encoding='utf-8')


def demo_supplier_details():
    """Пара контрагентов с дополненными реквизитами — чтобы было видно, как выглядит справочник."""
    r.save_supplier({'name': 'ООО "АЛЬФА"', 'inn': '0000000001', 'kpp': '770001001', 'ogrn': '1000000000001', 'smp': 'малое',
                     'address': 'г. Москва, ул. Примерная, д. 1', 'email': 'info@alfa.example'}, {'fio': 'Миранова В.С.'})


def izv_demo(today):
    """ОАО → «Подготовка извещения»: три закупки в работе (черновик, на согласовании, размещение)
    из вымышленных заявок (папка «Демо-данные (подготовка извещения)»), четвёртая заявка — для показа «+ Новая закупка»."""
    from openpyxl import load_workbook
    seeds = HERE / 'Демо-данные (подготовка извещения)'
    d = r.ПАПКА_ПАНЕЛИ / 'Подготовка извещения'
    shutil.rmtree(d, ignore_errors=True)
    sp = d / 'Справочники'
    shutil.copytree(r.ПАПКА_МОДУЛЯ_ИЗВЕЩЕНИЯ / 'dannye' / 'spravochniki', sp)
    wb = load_workbook(sp / 'Сотрудники отдела.xlsx')          # в демо — условные телефоны
    ws = wb['Сотрудники']
    for i, row in enumerate(ws.iter_rows(min_row=2), start=1):
        if row[0].value:
            row[1].value, row[2].value = str(2000 + i), '(495)000-00-00'
    wb.save(sp / 'Сотрудники отдела.xlsx')
    show = DOCS / '7 ОАО — подготовка извещения (перетащите в «+ Новая закупка»)'
    show.mkdir(parents=True, exist_ok=True)
    for f in seeds.glob('2155*.xlsx'):
        shutil.copy(f, show / f.name)
    r._IZV['key'] = None
    try:
        yadro, _pd, n = r.izv_core()
    except ValueError as e:
        print('  Подготовка извещения:', e)
        return
    folders = d / 'Папки закупок'

    def card(nomer, seed, who):
        f = folders / nomer
        f.mkdir(parents=True, exist_ok=True)
        src = next(seeds.glob(seed + '*.xlsx'))
        shutil.copy(src, f / ('Заявка ' + src.stem.split('заявка ')[-1] + '.xlsx'))
        st = yadro.otkryt_kartochku(n, f, who, nomer_zakupki=nomer)
        yadro.proverit(n, st['id'])
        return st['id']

    def to_soglasovanie(zid, plosh, days_ago):
        yadro.izmenit_pole(n, zid, 'ploshchadka', plosh)
        z = yadro._zagruzit(n, zid)
        forma = (yadro.sborka.podskazat_formu_dogovora(z['karta']['znacheniya'].get('predmet_im', ''),
                 yadro.sborka.zagruzit_formy_dogovora(n.sp(yadro.SPRAVOCHNIK_FORM_DOGOVORA))) or {}).get('nazvanie')
        yadro.izmenit_pole(n, zid, 'forma_dogovora', forma)
        yadro.podgotovit_prosmotr(n, zid, etap='soglasovanie')
        yadro.sohranit_prosmotr(n, zid)
        z = yadro._zagruzit(n, zid)
        z['data_etapa'] = (today - timedelta(days=days_ago)).isoformat()
        yadro._sohranit(n, zid, z)

    try:
        card('2140зэа', '2140', 'Гончарова Ксения Алексеевна')                     # черновик: решить площадку, форму договора, опечатку
        zb = card('2096зэа', '2096', 'Кузнецова Анна Станиславовна')
        to_soglasovanie(zb, 'https://etpgpb.ru', 5)                                 # на согласовании с …
        zd = card('2071зэа', '2071', 'Гончарова Ксения Алексеевна')
        to_soglasovanie(zd, 'http://223.astgoz.ru', 9)
        dr = yadro.data_razmeshcheniya_po_umolchaniyu(n)[0]
        z = yadro._zagruzit(n, zd)
        z['d_razmeshcheniya_vybrano'] = dr.isoformat()
        yadro._sohranit(n, zd, z)
        yadro.podgotovit_prosmotr(n, zd, etap='razmeshchenie', d_razmeshcheniya=dr)
        yadro.sohranit_prosmotr(n, zd)                                               # размещение: документ с датами готов
        for z in yadro.spisok_zakupok(n):
            yadro.osvobodit_kartochku(n, z['id'], None)
        print('  Подготовка извещения: 3 закупки в работе')
    except Exception as e:
        import traceback
        traceback.print_exc()
        print('  Подготовка извещения — демо не собралось:', e)


def build():
    print('Собираю демо-реестр…')
    shutil.rmtree(BASE, ignore_errors=True)
    r.ПАПКА_ПАНЕЛИ.mkdir(parents=True, exist_ok=True)
    demo_users()
    r.datetime = DemoDateTime
    random.seed(2026)
    today = date.today()
    at(date(today.year - 1, 1, 9), 9, 0)
    r.create_registry(REG(Y0), Y0)
    r.create_registry(REG(Y0 - 1), Y0 - 1)
    r.create_registry(REG(Y0 + 1), Y0 + 1)
    used = set()
    print(f'  {Y0 - 1}: весь год…')
    history(today, count=300, year=Y0 - 1, seed=1525, used=used)
    print(f'  {Y0}: с января…')
    history(today, used=used)
    seed_documents()
    plan = run_synthetic(today, synthetic(today))
    next_year(today, used)
    corrections(today, plan)
    legacy_names(plan)
    report_history(today)
    demo_supplier_details()
    protocol_folder(today, plan)
    analysts_register(today)
    demo_smp_categories()
    opaiz_demo(today)
    demo_extras(today)
    izv_demo(today)
    r.datetime = datetime
    import json
    j = r.load_journal()
    j.sort(key=lambda e: e.get('ts', ''))
    r.write_atomic(r.journal_path(), json.dumps(j, ensure_ascii=False, indent=1))
    # Нераспознанный файл во «Входящих» — чтобы было видно, как панель о нём сообщает.
    sample = next(DOCS.rglob('*договор*.docx'), None)
    if sample:
        inbox = r.year_obj(Y0).notices_in('nk')
        inbox.mkdir(parents=True, exist_ok=True)
        shutil.copy(sample, inbox / 'скан без шапки.docx')
    r.cache['mtime'] = None
    r.cache['years'] = {}
    print('Готово:', BASE)


if __name__ == '__main__':
    try:
        if '--оставить' not in sys.argv or not REG().exists():
            build()
        r.main()
    except KeyboardInterrupt:
        pass
    except Exception:
        import traceback
        traceback.print_exc()
        (HERE / 'ошибка запуска.txt').write_text(traceback.format_exc(), encoding='utf-8')
        print('\nНе запустилось. Текст ошибки сохранён в файл «ошибка запуска.txt» рядом с демо.py — пришлите его.')
        input('Нажмите Enter, чтобы закрыть окно…')
