# -*- coding: utf-8 -*-
"""Файлы старых форматов (раздел 9 ТЗ): .xls, .doc, .rtf.

Перед чтением программа молча переводит их в .xlsx / .docx во временной
папке программы; в папке закупки ничего не создаётся и не меняется. Весь
остальной код читает только .xlsx и .docx. Копия пересоздаётся, если
исходный файл изменился (поэтому «было → стало» работает и для них).

Основной способ — установленные Excel и Word: без окон, только чтение,
макросы выключены, ограничение по времени; зависший Excel/Word закрывается.
Запасной: .xls — библиотека xlrd; .doc/.rtf — текст и таблицы своим
разбором. Сотруднику о формате ничего не сообщается: файл выглядит как
обычный, факт перевода пишется только в журнал (zhurnal/sobytiya.log).
Сообщение сотруднику — только если файл защищён паролем или повреждён:
«Файл <имя> защищён паролем / повреждён — откройте его в Word/Excel и проверьте».
"""
from __future__ import annotations

import hashlib
import logging
import os
import re
import subprocess
import time
from pathlib import Path

STARYE_RASSHIRENIYA = {".xls", ".doc", ".rtf"}
OZHIDANIE_OFFICE = 120  # секунд на перевод одного файла
_BEZ_OKNA = getattr(subprocess, "CREATE_NO_WINDOW", 0)
_NET_PROGRAMMY = 3  # код выхода сценария: Excel/Word не установлен
# заведомо неверный пароль: у защищённого файла Excel/Word сразу отказывают,
# а не ждут ввода пароля в невидимом окне; у обычного файла пароль не мешает
_NEVERNYJ_PAROL = "podgotovka-izveshcheniya-net-parolya"
_CFB = bytes.fromhex("D0CF11E0A1B11AE1")  # контейнер, в котором лежат зашифрованные .xlsx/.docx

# факт перевода файла — только в журнал программы (zhurnal/sobytiya.log), не на экран
zhurnal = logging.getLogger("podgotovka_izveshcheniya.fajly")


class NeUdalosProchitat(Exception):
    """Файл защищён паролем или повреждён — текст для сотрудника."""

    def __init__(self, imya: str, parol: bool = False):
        programma = "Excel" if Path(imya).suffix.lower() in (".xls", ".xlsx") else "Word"
        prichina = "защищён паролем" if parol else "повреждён"
        super().__init__(f"Файл {imya} {prichina} — откройте его в {programma} и проверьте")
        self.parol = parol


def zashchishchen_parolem(put: Path) -> bool:
    """Защищён ли файл паролем (без Excel/Word)."""
    rasshirenie = put.suffix.lower()
    try:
        nachalo = put.read_bytes()[:8]
    except OSError:
        return False
    if rasshirenie in (".xlsx", ".docx"):
        return nachalo == _CFB  # обычный .xlsx/.docx — zip, зашифрованный — контейнер CFB
    if rasshirenie == ".xls":
        try:
            import xlrd
            xlrd.open_workbook(str(put), on_demand=True).release_resources()
        except Exception as e:
            return "encrypt" in str(e).lower()
    return False


def _papka_kopii(put: Path, papka_vremennyh: Path) -> Path:
    priznak = f"{put.resolve()}|{put.stat().st_mtime_ns}|{put.stat().st_size}"
    return Path(papka_vremennyh) / hashlib.sha1(priznak.encode("utf-8")).hexdigest()[:16]


def novaya_kopiya(put: str | Path, papka_vremennyh: str | Path) -> Path:
    """Путь к файлу в новом формате: для .xlsx/.docx — сам файл, для
    .xls/.doc/.rtf — временная копия (создаётся при первом обращении).
    Защищённый паролем или повреждённый файл — NeUdalosProchitat."""
    put = Path(put)
    rasshirenie = put.suffix.lower()
    if zashchishchen_parolem(put):
        raise NeUdalosProchitat(put.name, parol=True)
    if rasshirenie not in STARYE_RASSHIRENIYA:
        return put
    papka = _papka_kopii(put, Path(papka_vremennyh))
    kopiya = papka / (put.stem + (".xlsx" if rasshirenie == ".xls" else ".docx"))
    if kopiya.exists():
        return kopiya
    papka.mkdir(parents=True, exist_ok=True)
    vremennaya = kopiya.with_name("_" + kopiya.name)
    vremennaya.unlink(missing_ok=True)
    programma = "excel" if rasshirenie == ".xls" else "word"
    try:
        otvet = _cherez_office(put, vremennaya, programma)
        if otvet == "parol":
            raise NeUdalosProchitat(put.name, parol=True)
        sposob = "Excel" if programma == "excel" else "Word"
        if otvet != "ok":
            sposob = "xlrd" if programma == "excel" else "свой разбор текста"
            if programma == "excel":
                _xls_cherez_xlrd(put, vremennaya)
            else:
                _doc_rtf_svoim_razborom(put, vremennaya)
    except NeUdalosProchitat:
        raise
    except Exception:
        raise NeUdalosProchitat(put.name)
    if not vremennaya.exists():
        raise NeUdalosProchitat(put.name)
    vremennaya.replace(kopiya)
    zhurnal.info("файл %s прочитан через временную копию %s (%s)", put, kopiya.suffix, sposob)
    return kopiya


# ------------------------------------------------------------ Excel / Word

_SKRIPT = {
    "excel": (
        "$ErrorActionPreference = 'Stop';"
        "try { $p = New-Object -ComObject Excel.Application } catch { exit 3 };"
        "try { $p.Visible = $false; $p.DisplayAlerts = $false; $p.AskToUpdateLinks = $false;"
        "$p.AutomationSecurity = 3;"  # макросы не запускать
        # без обновления связей, только чтение, неверный пароль — без окна ввода пароля
        f"$k = $p.Workbooks.Open($env:PODG_IZ, 0, $true, 5, '{_NEVERNYJ_PAROL}');"
        "$k.SaveAs($env:PODG_V, 51); $k.Close($false) } finally { $p.Quit() }"
    ),
    "word": (
        "$ErrorActionPreference = 'Stop';"
        "try { $p = New-Object -ComObject Word.Application } catch { exit 3 };"
        "try { $p.Visible = $false; $p.DisplayAlerts = 0; $p.AutomationSecurity = 3;"
        # без вопросов, только чтение, неверный пароль — без окна ввода пароля
        f"$d = $p.Documents.Open($env:PODG_IZ, $false, $true, $false, '{_NEVERNYJ_PAROL}');"
        "$d.SaveAs2($env:PODG_V, 16); $d.Close($false) } finally { $p.Quit() }"
    ),
}
_PROCESS_OFFICE = {"excel": "EXCEL.EXE", "word": "WINWORD.EXE"}


def _processy(imya: str) -> set[int]:
    try:
        vyvod = subprocess.run(
            ["tasklist", "/FI", f"IMAGENAME eq {imya}", "/FO", "CSV", "/NH"],
            capture_output=True, stdin=subprocess.DEVNULL, timeout=20, creationflags=_BEZ_OKNA,
        ).stdout.decode("cp866", errors="ignore")
    except (OSError, subprocess.TimeoutExpired):
        return set()
    return {int(m.group(1)) for m in re.finditer(r'^"[^"]+","(\d+)"', vyvod, re.M)}


def _cherez_office(put: Path, kopiya: Path, programma: str) -> str:
    """Перевод через Excel/Word: «ok»; «parol» — файл защищён паролем;
    «net» — программы нет или она не справилась (тогда — запасной способ)."""
    do = _processy(_PROCESS_OFFICE[programma])
    okruzhenie = dict(os.environ, PODG_IZ=str(put.resolve()), PODG_V=str(kopiya.resolve()))
    otvet = "net"
    try:
        rezultat = subprocess.run(
            ["powershell", "-NoProfile", "-ExecutionPolicy", "Bypass", "-Command", _SKRIPT[programma]],
            capture_output=True, stdin=subprocess.DEVNULL, timeout=OZHIDANIE_OFFICE, env=okruzhenie,
            creationflags=_BEZ_OKNA,
        )
        if rezultat.returncode == 0 and kopiya.exists():
            otvet = "ok"
        elif re.search(r"парол|password", rezultat.stderr.decode("cp866", errors="ignore"), re.I):
            otvet = "parol"
    except (OSError, subprocess.TimeoutExpired):
        pass
    if otvet != "ok":
        # зависший Excel/Word, запущенный для перевода, — закрыть
        time.sleep(0.5)
        for pid in _processy(_PROCESS_OFFICE[programma]) - do:
            subprocess.run(["taskkill", "/PID", str(pid), "/F"], capture_output=True,
                           stdin=subprocess.DEVNULL, creationflags=_BEZ_OKNA)
        kopiya.unlink(missing_ok=True)
    return otvet


# ------------------------------------------------------------ запасные способы

def _xls_cherez_xlrd(put: Path, kopiya: Path) -> None:
    import openpyxl
    import xlrd

    try:
        kniga = xlrd.open_workbook(str(put))
    except Exception:
        raise NeUdalosProchitat(put.name)
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for list_ in kniga.sheets():
        ws = wb.create_sheet(list_.name[:31])
        for r in range(list_.nrows):
            for c in range(list_.ncols):
                tip = list_.cell_type(r, c)
                znachenie = list_.cell_value(r, c)
                if tip in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK, xlrd.XL_CELL_ERROR) or znachenie == "":
                    continue
                if tip == xlrd.XL_CELL_DATE:
                    try:
                        znachenie = xlrd.xldate_as_datetime(znachenie, kniga.datemode)
                    except Exception:
                        pass
                elif tip == xlrd.XL_CELL_NUMBER and float(znachenie).is_integer():
                    znachenie = int(znachenie)  # как Excel сохраняет целые в .xlsx
                elif tip == xlrd.XL_CELL_BOOLEAN:
                    znachenie = bool(znachenie)
                ws.cell(row=r + 1, column=c + 1, value=znachenie)
        for r1, r2, c1, c2 in list_.merged_cells:
            if r2 - r1 > 1 or c2 - c1 > 1:
                ws.merge_cells(start_row=r1 + 1, end_row=r2, start_column=c1 + 1, end_column=c2)
    wb.save(kopiya)


def _doc_rtf_svoim_razborom(put: Path, kopiya: Path) -> None:
    """Без Word: текст (и таблицы RTF) в .docx — для чтения хватает."""
    from docx import Document

    syroe = put.read_bytes()
    if syroe.lstrip().startswith(b"{\\rtf"):
        bloki = _razobrat_rtf(syroe.decode("latin-1"))
    else:
        bloki = [("p", s) for s in _tekst_iz_doc(syroe)]
    if not any(b[1] for b in bloki):
        raise NeUdalosProchitat(put.name)
    doc = Document()
    for vid, soderzhimoe in bloki:
        if vid == "t":
            stolbcov = max(len(r) for r in soderzhimoe)
            t = doc.add_table(rows=len(soderzhimoe), cols=stolbcov)
            for i, ryad in enumerate(soderzhimoe):
                for j, yach in enumerate(ryad):
                    t.cell(i, j).text = yach
        else:
            doc.add_paragraph(soderzhimoe)
    doc.save(kopiya)


def _tekst_iz_doc(syroe: bytes) -> list[str]:
    """Word 97–2003 без Word: текст лежит кусками UTF-16 — берём связные
    русско-латинские фрагменты (таблицы становятся строками)."""
    tekst = syroe.decode("utf-16-le", errors="ignore")
    kuski = re.findall(r"[\w\s.,:;!?()«»\"'%№/\\\-–—+=*]{12,}", tekst)
    stroki = []
    for kusok in kuski:
        for s in re.split(r"[\r\x07\x0b]+", kusok):
            s = s.strip()
            if len(s) >= 3 and re.search(r"[А-Яа-яЁё]", s):
                stroki.append(s)
    return stroki


def _razobrat_rtf(rtf: str) -> list[tuple[str, object]]:
    """Простой разбор RTF: абзацы и таблицы (\\cell, \\row); русские буквы —
    \\'xx (cp1251) и \\uN."""
    bloki: list[tuple[str, object]] = []
    abzac: list[str] = []
    yacheyka: list[str] = []
    ryad: list[str] = []
    tablica: list[list[str]] = []
    v_tablice = False
    glubina_propuska = []
    glubina = 0
    i = 0
    propustit_simvolov = 0
    sluzhebnye = ("fonttbl", "colortbl", "stylesheet", "info", "pict", "header", "footer", "*")

    def tekushchij() -> list[str]:
        return yacheyka if v_tablice else abzac

    while i < len(rtf):
        ch = rtf[i]
        if ch == "{":
            glubina += 1
            i += 1
            continue
        if ch == "}":
            if glubina_propuska and glubina_propuska[-1] == glubina:
                glubina_propuska.pop()
            glubina -= 1
            i += 1
            continue
        propusk = bool(glubina_propuska)
        if ch == "\\":
            m = re.match(r"\\([a-z]+)(-?\d+)? ?|\\'([0-9a-fA-F]{2})|\\(.)", rtf[i:])
            if not m:
                i += 1
                continue
            i += m.end()
            if m.group(3):
                if propustit_simvolov:
                    propustit_simvolov -= 1
                elif not propusk:
                    tekushchij().append(bytes([int(m.group(3), 16)]).decode("cp1251", errors="ignore"))
                continue
            if m.group(4):
                if m.group(4) == "*":
                    glubina_propuska.append(glubina)
                elif not propusk and m.group(4) in "\\{}":
                    tekushchij().append(m.group(4))
                continue
            slovo, chislo = m.group(1), m.group(2)
            if slovo in sluzhebnye and not propusk:
                glubina_propuska.append(glubina)
                continue
            if propusk:
                continue
            if slovo == "u" and chislo:
                kod = int(chislo)
                tekushchij().append(chr(kod + 65536 if kod < 0 else kod))
                propustit_simvolov = 1
            elif slovo == "intbl":
                v_tablice = True
            elif slovo == "cell":
                ryad.append("".join(yacheyka).strip())
                yacheyka = []
            elif slovo == "row":
                tablica.append(ryad)
                ryad = []
            elif slovo in ("par", "line"):
                if v_tablice:
                    yacheyka.append("\n")
                else:
                    bloki.append(("p", "".join(abzac).strip()))
                    abzac = []
            elif slovo == "pard":
                if v_tablice and tablica and not ryad and not yacheyka:
                    bloki.append(("t", tablica))
                    tablica = []
                v_tablice = False
            elif slovo == "tab":
                tekushchij().append("\t")
            continue
        if ch in "\r\n":
            i += 1
            continue
        if not propusk:
            if propustit_simvolov:
                propustit_simvolov -= 1
            else:
                tekushchij().append(ch)
        i += 1
    if tablica:
        bloki.append(("t", tablica))
    if abzac:
        bloki.append(("p", "".join(abzac).strip()))
    return [b for b in bloki if b[1]]
