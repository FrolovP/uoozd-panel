# -*- coding: utf-8 -*-
"""Роль файла в папке закупки (раздел 9 ТЗ): заявка, техническое задание,
расчёт НМЦК, инструкция по описанию товара, смета, не используется.
Определяется по содержимому, не по имени файла; сотрудник может сменить.
"""
from __future__ import annotations

import re
from pathlib import Path

import openpyxl
from docx import Document

SLUZHEBNYE_FAJLY = {"thumbs.db", "desktop.ini"}
# .xls/.doc/.rtf читаются через временную копию .xlsx/.docx (staryj_format)
RASSHireniya = {".xlsx", ".xls", ".docx", ".doc", ".rtf"}
# PDF собираются и показываются в списке файлов (роль по умолчанию —
# «не используется», программа их не читает без OCR-сервера), но не
# участвуют в опознавании роли по содержимому.
RASSHireniya_SPISOK = RASSHireniya | {".pdf"}


def sobrat_fajly_zakupki(papka: str | Path) -> list[Path]:
    """Файлы папки закупки, со всеми вложенными папками (пустые папки и
    папки-обёртки сами по себе файлов не дают — раздел 9 ТЗ, 14 ТЗ)."""
    rezultat = []
    for put in Path(papka).rglob("*"):
        if put.is_dir():
            continue
        if put.name.startswith("~$") or put.name.startswith("."):
            continue
        if put.name.lower() in SLUZHEBNYE_FAJLY:
            continue
        if put.suffix.lower() not in RASSHireniya_SPISOK:
            continue
        rezultat.append(put)
    return rezultat


def propushcheno_tekst(papka: str | Path) -> str:
    """Фраза «Пропущено: …» для экрана «Новая закупка» (раздел 9 ТЗ): сколько
    пустых папок и служебных файлов программа не показала в списке."""
    papka = Path(papka)
    pustye = 0
    for d in papka.rglob("*"):
        if d.is_dir() and not any(f.is_file() for f in d.rglob("*")):
            pustye += 1
    sluzhebnye = 0
    for f in papka.rglob("*"):
        if f.is_file() and (f.name.startswith("~$") or f.name.lower() in SLUZHEBNYE_FAJLY):
            sluzhebnye += 1
    kuski = []
    if pustye:
        kuski.append(f"{pustye} пуст{'ая папка' if pustye == 1 else 'ые папки и папки-обёртки'}")
    if sluzhebnye:
        kuski.append(f"{sluzhebnye} служебн{'ый файл' if sluzhebnye == 1 else 'ых файла'}")
    tekst = ", ".join(kuski) if kuski else "ничего"
    return f"Пропущено: {tekst}."


def _tekst_dlya_opoznaniya(put: Path) -> str:
    """Первые абзацы/ячейки файла — достаточно для поиска признаков роли."""
    try:
        if put.suffix.lower() not in (".docx", ".xlsx"):
            return ""
        if put.suffix.lower() == ".docx":
            doc = Document(put)
            kuski = [p.text for p in doc.paragraphs[:60]]
            for t in doc.tables[:2]:
                for row in t.rows[:3]:
                    kuski.extend(c.text for c in row.cells)
            return "\n".join(kuski)
        wb = openpyxl.load_workbook(put, data_only=True, read_only=True)
        kuski = []
        for ws in wb.worksheets:
            for row in ws.iter_rows(max_row=80, values_only=True):
                kuski.extend(str(v) for v in row if v)
        return "\n".join(kuski)
    except Exception:
        return ""


def dokument_programmy(imya: str) -> bool:
    """Готовый документ, который сделала сама программа («Аукцион. Извещение
    № …», «Аукцион. Проект договора № …») — не исходный файл закупки."""
    return bool(re.match(r"^Аукцион\.\s+(Извещение|Проект договора)\s+№", imya))


def rol_po_imeni(imya: str) -> str | None:
    """Файл со словом «заявка» в имени — по умолчанию «Заявка» («Заявка.xls»,
    «Заявка КЦ 1500-а-26.docx»; но не «ТЗ (приложение 2 к заявке)»)."""
    if re.search(r"(?<![а-яё])заявка(?![а-яё])", Path(imya).stem.lower()):
        return "zayavka"
    return None


def opredelit_rol(put: str | Path, tekst: str | None = None, imya: str | None = None) -> str:
    """Роль по содержимому (раздел 9 ТЗ): «zayavka», «tz», «raschet_nmck»,
    «instrukciya», «smeta» или «ne_ispolzuetsya». put — файл для чтения
    (.xlsx/.docx, для старых форматов — временная копия); imya — имя файла
    сотрудника, если отличается."""
    put = Path(put)
    imya = imya or put.name
    if imya.startswith("~$") or imya.lower() in SLUZHEBNYE_FAJLY:
        return "ne_ispolzuetsya"
    if Path(imya).suffix.lower() not in RASSHireniya or dokument_programmy(imya):
        return "ne_ispolzuetsya"

    nizhnij = (tekst if tekst is not None else _tekst_dlya_opoznaniya(put)).lower()

    if "заявка на закрытый аукцион" in nizhnij or "лист согласования к заявке" in nizhnij:
        return "zayavka"
    if rol_po_imeni(imya):
        return "zayavka"
    if "смета" in nizhnij or "локальный сметный расчёт" in nizhnij:
        return "smeta"
    if (
        "используемый метод определения нмцк" in nizhnij
        or "цены поставщиков (коммерческие предложения)" in nizhnij
        or "расчет нмцк" in nizhnij or "расчёт нмцк" in nizhnij
    ):
        return "raschet_nmck"
    # у некоторых заявок один файл ТЗ содержит и инструкцию по описанию
    # товара (раздел 9 ТЗ, пример 426зэа) — техническое задание первично
    if "техническое задание" in nizhnij or "требуемые функции" in nizhnij or "технические характеристики" in nizhnij:
        return "tz"
    if "инструкция по описанию товара" in nizhnij:
        return "instrukciya"
    return "ne_ispolzuetsya"
