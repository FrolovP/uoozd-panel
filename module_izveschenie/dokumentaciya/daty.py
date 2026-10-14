# -*- coding: utf-8 -*-
"""Даты извещения (раздел 8 ТЗ).

Производственный календарь — справочник Excel
(dannye/spravochniki/Производственный календарь.xlsx), который правит
специалист, без изменения кода. Лист «Исключения» перечисляет только дни,
отличающиеся от обычной пятидневной недели: «нерабочий» (праздник или
перенесённый выходной на будний день) и «рабочий» (рабочая суббота/воскресенье
по переносу). Обычные дни модуль определяет сам.
"""
from __future__ import annotations

from dataclasses import dataclass
from datetime import date, timedelta
from pathlib import Path

import openpyxl


@dataclass
class ProizvodstvennyjKalendar:
    iskhlyucheniya: dict[date, bool]  # True = рабочий, False = нерабочий (только отличия от обычной недели)
    gody: set[int]  # годы, на которые есть данные (лист «Проверка»)

    def rabochij_den(self, d: date) -> bool:
        if d in self.iskhlyucheniya:
            return self.iskhlyucheniya[d]
        return d.weekday() < 5

    def god_izvesten(self, god: int) -> bool:
        return god in self.gody

    def rabochih_dnej_v_mesyace(self, god: int, mesyac: int) -> int:
        d = date(god, mesyac, 1)
        schetchik = 0
        while d.month == mesyac:
            if self.rabochij_den(d):
                schetchik += 1
            d += timedelta(days=1)
        return schetchik


def zagruzit_kalendar(put: str | Path) -> ProizvodstvennyjKalendar:
    wb = openpyxl.load_workbook(put, data_only=True)
    ws_iskl = wb["Исключения"]
    iskhlyucheniya: dict[date, bool] = {}
    for row in ws_iskl.iter_rows(min_row=2, values_only=True):
        if not row or row[0] is None:
            continue
        d = row[0].date() if hasattr(row[0], "date") else row[0]
        tip = str(row[1]).strip().lower() if row[1] is not None else ""
        iskhlyucheniya[d] = (tip == "рабочий")

    ws_proverka = wb["Проверка"]
    gody = {
        int(row[0]) for row in ws_proverka.iter_rows(min_row=2, values_only=True)
        if row and isinstance(row[0], (int, float))
    }
    return ProizvodstvennyjKalendar(iskhlyucheniya, gody)


class OshibkaDat(Exception):
    """Самопроверка дат (раздел 11 ТЗ): даты не считаются, если не хватает
    данных или сотрудник выбрал неподходящий день."""


class NetKalendaryaNaGod(OshibkaDat):
    def __init__(self, god: int):
        super().__init__(f"нет производственного календаря на {god}")
        self.god = god


class DataRazmeshcheniyaNerabochaya(OshibkaDat):
    def __init__(self, d: date):
        super().__init__(f"{d:%d.%m.%Y} — нерабочий день. Выберите рабочий день")
        self.data = d


def _sleduyuschij_rabochij(d: date, kal: ProizvodstvennyjKalendar) -> date:
    d = d + timedelta(days=1)
    while not kal.rabochij_den(d):
        d += timedelta(days=1)
    return d


def _tot_zhe_ili_sleduyuschij_rabochij(d: date, kal: ProizvodstvennyjKalendar) -> date:
    while not kal.rabochij_den(d):
        d += timedelta(days=1)
    return d


def _otnyat_rabochie_dni(d: date, n: int, kal: ProizvodstvennyjKalendar) -> date:
    while n > 0:
        d -= timedelta(days=1)
        if kal.rabochij_den(d):
            n -= 1
    return d


@dataclass
class DatyIzveshcheniya:
    d_razmeshcheniya: date
    d_zaprosy_do: date
    d_podacha_do: date
    d_rassmotrenie: date
    d_aukcion: date
    d_itogi: date


def rasschitat_daty(d_razmeshcheniya: date, kal: ProizvodstvennyjKalendar) -> DatyIzveshcheniya:
    """Раздел 8 ТЗ. Окончание подачи — 15 полных дней со следующего дня после
    размещения, затем первый рабочий день; запросы — до четвёртого рабочего
    дня перед окончанием подачи; рассмотрение, аукцион, итоги — каждое на
    следующий рабочий день."""
    if not kal.god_izvesten(d_razmeshcheniya.year):
        raise NetKalendaryaNaGod(d_razmeshcheniya.year)
    if not kal.rabochij_den(d_razmeshcheniya):
        raise DataRazmeshcheniyaNerabochaya(d_razmeshcheniya)

    # следующий день после размещения + 15 полных дней = +16 календарных дней
    baza_podachi = d_razmeshcheniya + timedelta(days=16)
    d_podacha_do = _tot_zhe_ili_sleduyuschij_rabochij(baza_podachi, kal)
    d_zaprosy_do = _otnyat_rabochie_dni(d_podacha_do, 4, kal)
    d_rassmotrenie = _sleduyuschij_rabochij(d_podacha_do, kal)
    d_aukcion = _sleduyuschij_rabochij(d_rassmotrenie, kal)
    d_itogi = _sleduyuschij_rabochij(d_aukcion, kal)

    for d in (d_zaprosy_do, d_podacha_do, d_rassmotrenie, d_aukcion, d_itogi):
        if not kal.god_izvesten(d.year):
            raise NetKalendaryaNaGod(d.year)

    return DatyIzveshcheniya(d_razmeshcheniya, d_zaprosy_do, d_podacha_do, d_rassmotrenie, d_aukcion, d_itogi)


@dataclass
class OstalnyeDaty:
    d_zaprosy_do: date
    d_rassmotrenie: date
    d_aukcion: date
    d_itogi: date


def rasschitat_ostalnye_daty(d_podacha_do: date, kal: ProizvodstvennyjKalendar) -> OstalnyeDaty:
    """От выбранной даты окончания подачи (раздел 8 ТЗ, экран «Календарь») —
    запросы разъяснений (до четвёртого рабочего дня перед окончанием подачи),
    рассмотрение/аукцион/итоги — каждое на следующий рабочий день. Не зависит
    от даты размещения: сотрудник может выбрать окончание подачи позже
    минимально допустимой даты, остальные даты считаются от неё."""
    return OstalnyeDaty(
        d_zaprosy_do=_otnyat_rabochie_dni(d_podacha_do, 4, kal),
        d_rassmotrenie=_sleduyuschij_rabochij(d_podacha_do, kal),
        d_aukcion=_sleduyuschij_rabochij(_sleduyuschij_rabochij(d_podacha_do, kal), kal),
        d_itogi=_sleduyuschij_rabochij(_sleduyuschij_rabochij(_sleduyuschij_rabochij(d_podacha_do, kal), kal), kal),
    )


def format_data(d: date) -> str:
    return f"{d:%d.%m.%Y}г."


def format_vremya(vremya_chch_mm: str) -> str:
    """«09:00» -> «09 ч. 00 мин.» — вид для документа."""
    chasy, minuty = vremya_chch_mm.split(":")
    return f"{int(chasy):02d} ч. {int(minuty):02d} мин."


def format_dannye_dlya_shablona(daty: DatyIzveshcheniya, t_podacha: str = "09:00", t_aukcion: str = "11:00") -> dict[str, str]:
    """Поля карточки в виде, готовом для шаблона (раздел 5 ТЗ): время подачи
    и время аукциона по умолчанию 09:00/11:00, сотрудник может поправить на
    экране «Размещение» (решение от 11.10.2026)."""
    return {
        "d_razmeshcheniya": format_data(daty.d_razmeshcheniya),
        "d_zaprosy_do": format_data(daty.d_zaprosy_do),
        "d_podacha_do": format_data(daty.d_podacha_do),
        "t_podacha_do": format_vremya(t_podacha),
        "d_rassmotrenie": format_data(daty.d_rassmotrenie),
        "d_aukcion": format_data(daty.d_aukcion),
        "t_aukcion": format_vremya(t_aukcion),
        "d_itogi": format_data(daty.d_itogi),
    }
