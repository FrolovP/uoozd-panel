# -*- coding: utf-8 -*-
"""Сборка извещения и проекта договора (разделы 5 и 8 ТЗ).

Документ собирается из карточки с включённым экранированием (autoescape):
в тексте заявок встречаются знаки «<», «>» и «&» (раздел 5 ТЗ). Таблица ТЗ
и расчёт НМЦК «по позициям»/тарифным методом — вложенные документы
(docxcompose, через docxtpl[subdoc]).
"""
from __future__ import annotations

from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Twips
from docxtpl import DocxTemplate

import tablica_tz


def _postavit_v_yacheyku(cell, text: str, shrift: str, razmer: int, bold: bool = False, left: bool = False) -> None:
    cell.text = ""
    for i, line in enumerate(str(text).split("\n")):
        p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT if left else WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(line.strip())
        r.bold = bold
        r.font.name = shrift
        r.font.size = Pt(razmer)


def sobrat_tz_docx(zagolovki: list[str], stroki: list[list[str]], put: str | Path,
                   fajly_tz: list[Path] | None = None, forma_426: bool = False,
                   ispravleniya: dict[str, str] | None = None) -> None:
    """Таблица ТЗ как отдельный документ (вложенный в извещение
    через tpl.new_subdoc) — оформление как в tz_iz_excel.py. fajly_tz —
    отдельные файлы ТЗ: их таблицы (характеристики) идут следом целиком."""
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Times New Roman"
    st.font.size = Pt(9)

    shirina_kolonki = tablica_tz.SHIRINY_STOLBCOV
    ostalnoe = 14200 - sum(shirina_kolonki.get(n, 0) for n in zagolovki)
    bez_shiriny = sum(1 for n in zagolovki if n not in shirina_kolonki) or 1
    shiriny = [shirina_kolonki.get(n) or ostalnoe // bez_shiriny for n in zagolovki]

    t = doc.add_table(rows=0, cols=len(zagolovki))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER

    row = t.add_row()
    for i, n in enumerate(zagolovki):
        row.cells[i].width = Twips(shiriny[i])
        _postavit_v_yacheyku(row.cells[i], n, "Times New Roman", 9, bold=True)
    trPr = row._tr.get_or_add_trPr()
    h = OxmlElement("w:tblHeader")
    trPr.append(h)

    for stroka in stroki:
        row = t.add_row()
        for i, znachenie in enumerate(stroka):
            row.cells[i].width = Twips(shiriny[i])
            sleva = zagolovki[i].lower().startswith("техническ")
            _postavit_v_yacheyku(row.cells[i], znachenie, "Times New Roman", 9, left=sleva)

    t.autofit = False
    grid = t._tbl.tblGrid
    for gc, w in zip(grid.findall(qn("w:gridCol")), shiriny):
        gc.set(qn("w:w"), str(w))
    tblPr = t._tbl.tblPr
    tw = tblPr.find(qn("w:tblW"))
    if tw is None:
        tw = OxmlElement("w:tblW")
        tblPr.append(tw)
    tw.set(qn("w:w"), str(sum(shiriny)))
    tw.set(qn("w:type"), "dxa")
    lay = OxmlElement("w:tblLayout")
    lay.set(qn("w:type"), "fixed")
    tblPr.append(lay)
    for fajl in fajly_tz or []:
        for tablica in tablicy_iz_fajla_tz(fajl):
            doc.add_paragraph()
            vstavit_tablicu(doc, tablica, forma_426, ispravleniya)
    doc.save(put)


def tablicy_iz_fajla_tz(put: str | Path) -> list:
    """Таблицы характеристик из отдельного файла ТЗ (.docx): все таблицы файла."""
    return list(Document(put).tables)


def vstavit_tablicu(doc, tablica, forma_426: bool = False, ispravleniya: dict[str, str] | None = None) -> None:
    """Таблица из другого файла — целиком, с оформлением; названия столбцов
    шапки — к виду документации (список замен tablica_tz); опечатки, которые
    сотрудник исправил («Исправить»), исправляются и здесь."""
    import copy

    tbl = copy.deepcopy(tablica._tbl)
    doc.element.body.append(tbl)
    from docx.table import Table
    nova = Table(tbl, doc)
    stil = tbl.tblPr.find(qn("w:tblStyle")) if tbl.tblPr is not None else None
    if stil is not None and stil.get(qn("w:val")) not in {s.style_id for s in doc.styles}:
        nova.style = "Table Grid"
    elif stil is None and tbl.tblPr.find(qn("w:tblBorders")) is None:
        nova.style = "Table Grid"
    import re

    for najdeno, ispravlenie in (ispravleniya or {}).items():
        obrazec = re.compile(rf"(?<!\w){re.escape(najdeno)}(?!\w)")
        for t_ in tbl.iter(qn("w:t")):
            if t_.text and obrazec.search(t_.text):
                t_.text = obrazec.sub(lambda _: ispravlenie, t_.text)
    if nova.rows:
        for cell in nova.rows[0].cells:
            staroe = cell.text.strip()
            novoe = tablica_tz.pereimenovat_zagolovok(staroe, forma_426)
            if novoe != staroe:
                _zamenit_tekst_yacheyki(cell, novoe)


def _zamenit_tekst_yacheyki(cell, tekst: str) -> None:
    """Новый текст ячейки с оформлением её первого фрагмента."""
    pervyj = True
    for p in cell.paragraphs:
        for r in p.runs:
            r.text = tekst if pervyj else ""
            pervyj = False
    if pervyj:
        cell.paragraphs[0].add_run(tekst)


def sobrat_nmck_tarif_docx(nmck_tarif: dict[str, Any], put: str | Path) -> None:
    """Приложение № 2 при тарифном методе — как в
    shablon/skripty_dlya_spravki/nmck_tarif.py, но из уже извлечённых данных
    карточки, а не из файла заявки заново (раздел 12 ТЗ)."""
    pozicii = nmck_tarif["pozicii"]
    nazvaniya = list(pozicii[0].keys())

    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Arial"
    st.font.size = Pt(10)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    st.paragraph_format.space_after = Pt(0)
    widths = [700, 2000, 1700, 1800, 1000, 2140]  # всего 9340, как в шаблоне
    t = doc.add_table(rows=0, cols=6)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER

    def add(cells, bold=False, align=None):
        row = t.add_row()
        for i, w in enumerate(widths):
            row.cells[i].width = Twips(w)
        pos = 0
        for text, span in cells:
            c = row.cells[pos]
            if span > 1:
                c = c.merge(row.cells[pos + span - 1])
            c.text = ""
            p = c.paragraphs[0]
            run = p.add_run(str(text))
            run.bold = bold
            run.font.name = "Arial"
            run.font.size = Pt(10)
            run._r.get_or_add_rPr().get_or_add_rFonts().set(qn("w:cs"), "Arial")
            if align:
                p.alignment = align
            pos += span
        return row

    C = WD_ALIGN_PARAGRAPH.CENTER
    add([("Основные характеристики объекта закупки (работ, услуг)", 2), (nmck_tarif["obekt"], 4)])
    add([("Используемый метод определения НМЦК с обоснованием:", 2), (nmck_tarif["metod"], 4)])
    add([("Расчет НМЦК: " + nmck_tarif["reestr_tekst"], 6)])
    add([("№", 1), ("Цена по реестру", 1), (nmck_tarif["nds_zagolovok"], 1),
         ("Торговая надбавка", 1), ("Кол-во", 1), ("Сумма", 1)], bold=True, align=C)
    for poz in pozicii:
        znacheniya = [poz[n] for n in nazvaniya]
        add([(znacheniya[0], 1)] + [(x, 1) for x in znacheniya[1:]], align=C)
    add([("ИТОГО:", 5), (nmck_tarif["itogo"], 1)], bold=True, align=C)
    for row in t.rows:
        trPr = row._tr.get_or_add_trPr()
        cs = OxmlElement("w:cantSplit")
        trPr.append(cs)
    doc.save(put)


def sobrat_nmck_po_poziciyam_docx(nmck_poz: dict[str, Any], put: str | Path) -> None:
    """Приложение № 2 при расчёте «по позициям» — таблица листа «Расчет НМЦК»
    заявки как есть (образец — документация 426зэа): шапка в две строки
    (поставщики с входящими номерами КП), строки позиций, ИТОГО."""
    verh, niz = nmck_poz["verh"], nmck_poz.get("niz")
    n = len(verh)
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Times New Roman"
    st.font.size = Pt(9)
    st.paragraph_format.space_after = Pt(0)
    t = doc.add_table(rows=0, cols=n)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER

    def stroka(znacheniya, bold=False):
        row = t.add_row()
        for i, z in enumerate(znacheniya):
            _postavit_v_yacheyku(row.cells[i], z, "Times New Roman", 9, bold=bold, left=(i == 1 and not bold))
        return row

    r1 = stroka(verh, bold=True)
    if niz:
        r2 = stroka(niz, bold=True)
        i = 0
        while i < n:
            j = i
            while j + 1 < n and not verh[j + 1] and niz[j + 1]:
                j += 1
            if j > i:  # «№ коммерческого предложения» над столбцами поставщиков
                r1.cells[i].merge(r1.cells[j])
                _postavit_v_yacheyku(r1.cells[i], verh[i], "Times New Roman", 9, bold=True)
            elif i > 0 and verh[i] and not niz[i]:
                r1.cells[i].merge(r2.cells[i])
                _postavit_v_yacheyku(r1.cells[i], verh[i], "Times New Roman", 9, bold=True)
            i = j + 1
    for s in nmck_poz["stroki"]:
        stroka(s)
    if nmck_poz.get("itogo"):
        c = nmck_poz.get("itogo_stolbec") or n - 1
        row = t.add_row()
        if c > 0:
            row.cells[0].merge(row.cells[c - 1])
        row.cells[c].merge(row.cells[n - 1])
        _postavit_v_yacheyku(row.cells[0], "ИТОГО:", "Times New Roman", 9, bold=True)
        _postavit_v_yacheyku(row.cells[c], nmck_poz["itogo"], "Times New Roman", 9, bold=True)
    doc.save(put)


def dannye_dlya_shablona(znacheniya: dict[str, Any]) -> dict[str, Any]:
    """Карточка -> данные для рендера: служебные поля (с «_» в начале имени —
    подсказки, не для шаблона) убираются, как в proverka_shablona.py."""
    return {k: v for k, v in znacheniya.items() if not k.startswith("_")}


def _nuzhna_tablica_nmck(nmck_raschet: dict[str, Any]) -> bool:
    if nmck_raschet.get("smeta"):
        return False
    if nmck_raschet.get("po_obshchej_summe") and not nmck_raschet.get("po_minimalnomu"):
        return False
    return True


def sobrat_izveshchenie(
    put_shablona: str | Path,
    znacheniya: dict[str, Any],
    pozicii: list[dict],
    put_rezultata: str | Path,
    put_vremennoj_papki: str | Path,
    tablica_tz_gotovaya: tuple[list[str], list[list[str]]] | None = None,
    fajly_tz: list[Path] | None = None,
    ispravleniya: dict[str, str] | None = None,
) -> None:
    """Собирает извещение из шаблона и карточки (разделы 5, 8 ТЗ).
    tablica_tz_gotovaya — таблица ТЗ, уже настроенная сотрудником на экране
    «Таблица ТЗ» (порядок, скрытые/возвращённые/добавленные столбцы); если не
    передана, таблица собирается заново из позиций заявки без правок."""
    tpl = DocxTemplate(put_shablona)
    dannye = dannye_dlya_shablona(znacheniya)
    vremennaya = Path(put_vremennoj_papki)
    vremennaya.mkdir(parents=True, exist_ok=True)

    if tablica_tz_gotovaya is not None:
        zagolovki, stroki = tablica_tz_gotovaya
    else:
        zagolovki, stroki = tablica_tz.sobrat_tablicu_tz(pozicii)
    put_tz = vremennaya / "_tz_tablica.docx"
    stolbcy_zayavki = list(pozicii[0].keys()) if pozicii else []
    sobrat_tz_docx(zagolovki, stroki, put_tz, fajly_tz, tablica_tz.eto_forma_426(stolbcy_zayavki), ispravleniya)
    dannye["tz_tablicy"] = tpl.new_subdoc(put_tz)

    nmck_raschet = dict(dannye.get("nmck_raschet") or {})
    if _nuzhna_tablica_nmck(nmck_raschet):
        nmck_tarif = znacheniya.get("_nmck_tarif")
        nmck_poz = znacheniya.get("_nmck_poz")
        if nmck_tarif:
            put_tarif = vremennaya / "_nmck_tablica.docx"
            sobrat_nmck_tarif_docx(nmck_tarif, put_tarif)
            nmck_raschet["tablica"] = tpl.new_subdoc(put_tarif)
            dannye["nmck_raschet"] = nmck_raschet
        elif nmck_poz:
            put_poz = vremennaya / "_nmck_tablica.docx"
            sobrat_nmck_po_poziciyam_docx(nmck_poz, put_poz)
            nmck_raschet["tablica"] = tpl.new_subdoc(put_poz)
            dannye["nmck_raschet"] = nmck_raschet

    _pustye_dlya_otsutstvuyushchih(tpl, dannye)
    tpl.render(dannye, autoescape=True)
    tpl.save(put_rezultata)


def _pustye_dlya_otsutstvuyushchih(tpl: DocxTemplate, dannye: dict[str, Any]) -> None:
    """Поле шаблона, которого нет в карточке, печатается пустым, а не роняет
    сборку; обязательные поля заранее проверяет самопроверка (список «Решить»)."""
    for imya in tpl.get_undeclared_template_variables():
        if imya not in dannye:
            dannye[imya] = ""


# ------------------------------------------------------------- проект договора

def zagruzit_formy_dogovora(put: str | Path) -> list[dict[str, str]]:
    """Справочник типовых форм договора — dannye/spravochniki/Формы договора.xlsx
    (столбцы: Название, Файл, Признаки, Форма-заглушка). Признаки — слова, по
    которым программа подсказывает форму; выбирает сотрудник (раздел 8 ТЗ).
    «Форма-заглушка» = «да» — форма временная, пока нет настоящей."""
    import openpyxl

    wb = openpyxl.load_workbook(put, data_only=True)
    ws = wb["Формы"]
    formy = []
    for nazvanie, fajl, priznaki, zaglushka in ws.iter_rows(min_row=2, max_col=4, values_only=True):
        if not nazvanie:
            continue
        formy.append({
            "nazvanie": str(nazvanie).strip(),
            "fajl": str(fajl).strip(),
            "priznaki": [p.strip().lower() for p in str(priznaki or "").split(",") if p.strip()],
            "zaglushka": str(zaglushka or "").strip().lower() == "да",
        })
    return formy


def podskazat_formu_dogovora(predmet_im: str, formy: list[dict[str, str]]) -> dict[str, str]:
    """Подсказывает форму по словам-признакам из предмета закупки; выбирает
    сотрудник (раздел 8 ТЗ). Первая форма без признаков — запасной вариант."""
    nizhnij = predmet_im.lower()
    zapasnaya = None
    for forma in formy:
        if not forma["priznaki"]:
            zapasnaya = zapasnaya or forma
            continue
        if any(priznak in nizhnij for priznak in forma["priznaki"]):
            return forma
    if zapasnaya is None and formy:
        zapasnaya = formy[0]
    return zapasnaya or {}


POMETKA_OBRAZEC = "ОБРАЗЕЦ — форма в разработке"


def sobrat_proekt_dogovora(
    put_shablona_formy: str | Path, znacheniya: dict[str, Any], put_rezultata: str | Path, zaglushka: bool = False,
) -> None:
    """Проект договора — отдельный файл по типовой форме (раздел 8 ТЗ).
    Текст формы не меняется, заполняются только пропуски из карточки:
    поставщик, цена победителя и номер договора — по результатам закупки,
    остаются как в форме. zaglushka — форма временная: первой строкой
    крупно «ОБРАЗЕЦ — форма в разработке»."""
    tpl = DocxTemplate(put_shablona_formy)
    dannye = dannye_dlya_shablona(znacheniya)
    _pustye_dlya_otsutstvuyushchih(tpl, dannye)
    tpl.render(dannye, autoescape=True)
    if zaglushka:
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Pt, RGBColor

        abzac = tpl.docx.add_paragraph()
        abzac.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = abzac.add_run(POMETKA_OBRAZEC)
        run.bold = True
        run.font.size = Pt(28)
        run.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
        telo = tpl.docx.element.body
        telo.remove(abzac._p)
        telo.insert(0, abzac._p)
    tpl.save(put_rezultata)
