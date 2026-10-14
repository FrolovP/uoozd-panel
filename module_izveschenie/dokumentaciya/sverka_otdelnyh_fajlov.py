# -*- coding: utf-8 -*-
"""Сверка отдельных файлов заявки (раздел 9 ТЗ): ТЗ, расчёт НМЦК, инструкция
по описанию товара в отдельном файле — сверяются с заявкой. При расхождении
ничего не берётся ни из заявки, ни из файла: возвращается список расхождений
«что — в заявке — в файле». Без расхождений список пуст, экран не показывается.
"""
from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path

from docx import Document


@dataclass
class Rashozhdenie:
    chto: str
    v_zayavke: str
    v_fajle: str

    def __str__(self) -> str:
        return f"{self.chto} — в заявке: «{self.v_zayavke}», в файле: «{self.v_fajle}»"


def sverit_nmck_kp(kp_iz_zayavki: list[dict], kp_iz_fajla: list[dict]) -> list[Rashozhdenie]:
    """Сверяет расчёт НМЦК «по общей сумме» по входящему номеру КП: дата и
    сумма должны совпасть (раздел 9 ТЗ). Если хоть одно КП разошлось —
    пересчитывается и сверяется средняя арифметическая по каждой стороне."""
    rashozhdeniya: list[Rashozhdenie] = []
    po_nomeru_z = {k["vh_nomer"]: k for k in kp_iz_zayavki}
    po_nomeru_f = {k["vh_nomer"]: k for k in kp_iz_fajla}

    for nomer in po_nomeru_z.keys() | po_nomeru_f.keys():
        z = po_nomeru_z.get(nomer)
        f = po_nomeru_f.get(nomer)
        if z is None:
            rashozhdeniya.append(Rashozhdenie(f"КП № {nomer}", "нет", f"есть, {f['summa']}"))
            continue
        if f is None:
            rashozhdeniya.append(Rashozhdenie(f"КП № {nomer}", f"есть, {z['summa']}", "нет"))
            continue
        if z["summa"] != f["summa"]:
            rashozhdeniya.append(Rashozhdenie(f"КП № {nomer}, сумма", z["summa"], f["summa"]))
        if z.get("vh_data") != f.get("vh_data"):
            rashozhdeniya.append(Rashozhdenie(f"КП № {nomer}, дата", z.get("vh_data", ""), f.get("vh_data", "")))

    if rashozhdeniya and kp_iz_zayavki and kp_iz_fajla:
        import raschety
        from izvlechenie_excel import format_chislo
        srednyaya_z = raschety.rasschitat_nmck_po_kp([raschety.chislo(k["summa"]) for k in kp_iz_zayavki]).srednyaya
        srednyaya_f = raschety.rasschitat_nmck_po_kp([raschety.chislo(k["summa"]) for k in kp_iz_fajla]).srednyaya
        if srednyaya_z != srednyaya_f:
            rashozhdeniya.append(Rashozhdenie("Средняя арифметическая", format_chislo(srednyaya_z), format_chislo(srednyaya_f)))
    return rashozhdeniya


def _normalizovat_abzac(s: str) -> str:
    return " ".join(s.split())


def sverit_instrukciyu(iz_zayavki: list[str], iz_fajla: list[str]) -> list[Rashozhdenie]:
    """Сверяет инструкцию по абзацам, без учёта пустых строк и лишних
    пробелов (раздел 9 ТЗ)."""
    a = [_normalizovat_abzac(s) for s in iz_zayavki if s.strip()]
    b = [_normalizovat_abzac(s) for s in iz_fajla if s.strip()]
    rashozhdeniya = []
    for i, (x, y) in enumerate(zip(a, b), start=1):
        if x != y:
            rashozhdeniya.append(Rashozhdenie(f"абзац {i}", x, y))
    if len(a) != len(b):
        rashozhdeniya.append(Rashozhdenie("число абзацев", str(len(a)), str(len(b))))
    return rashozhdeniya


def _stolbec(zagolovki: list[str], *klyuchi: str) -> int | None:
    for i, z in enumerate(zagolovki):
        nizhnij = " ".join(z.lower().split())
        if any(k in nizhnij for k in klyuchi):
            return i
    return None


def _pervaya_stroka(s: str) -> str:
    return " ".join((s or "").strip().split("\n")[0].split())


def procitat_pozicii_tz(put: str | Path) -> list[dict]:
    """Позиции из отдельного файла ТЗ (Word): первая таблица со столбцами
    «Наименование» и «Кол-во» — наименование, ед. изм., количество."""
    doc = Document(put)
    for t in doc.tables:
        if not t.rows:
            continue
        zagolovki = [c.text for c in t.rows[0].cells]
        i_naim = _stolbec(zagolovki, "наименование")
        i_kol = _stolbec(zagolovki, "кол-во", "количество")
        if i_naim is None or i_kol is None:
            continue
        i_ed = _stolbec(zagolovki, "ед. изм", "ед.изм", "единица")
        pozicii = []
        for row in t.rows[1:]:
            yach = [c.text for c in row.cells]
            if not _pervaya_stroka(yach[i_naim]):
                continue
            pozicii.append({
                "naimenovanie": _pervaya_stroka(yach[i_naim]),
                "ed": _pervaya_stroka(yach[i_ed]) if i_ed is not None else "",
                "kolichestvo": _pervaya_stroka(yach[i_kol]),
            })
        return pozicii
    return []


def pozicii_zayavki_dlya_sverki(pozicii: list[dict]) -> list[dict]:
    """Те же три поля из позиций карточки (столбцы заявки называются по-разному)."""
    rezultat = []
    for p in pozicii:
        zagolovki = list(p.keys())

        def znach(*klyuchi: str) -> str:
            i = _stolbec(zagolovki, *klyuchi)
            return _pervaya_stroka(str(p[zagolovki[i]].get("znachenie") or "")) if i is not None else ""

        rezultat.append({"naimenovanie": znach("наименование"), "ed": znach("ед. изм", "ед.изм", "единица"),
                         "kolichestvo": znach("кол-во", "количество")})
    return rezultat


def _odno_i_to_zhe(a: str, b: str) -> bool:
    a, b = a.lower().rstrip(" ."), b.lower().rstrip(" .")
    return a == b or a.startswith(b) or b.startswith(a)


def sverit_tz(iz_zayavki: list[dict], iz_fajla: list[dict]) -> list[Rashozhdenie]:
    """Позиции ТЗ из отдельного файла и из заявки: число позиций,
    наименование (первая строка), ед. изм., количество (раздел 9 ТЗ)."""
    rashozhdeniya = []
    if len(iz_zayavki) != len(iz_fajla):
        rashozhdeniya.append(Rashozhdenie("число позиций", str(len(iz_zayavki)), str(len(iz_fajla))))
    for n, (z, f) in enumerate(zip(iz_zayavki, iz_fajla), start=1):
        if not _odno_i_to_zhe(z["naimenovanie"], f["naimenovanie"]):
            rashozhdeniya.append(Rashozhdenie(f"позиция {n}, наименование", z["naimenovanie"], f["naimenovanie"]))
        if z["ed"] and f["ed"] and z["ed"].lower().rstrip(".") != f["ed"].lower().rstrip("."):
            rashozhdeniya.append(Rashozhdenie(f"позиция {n}, ед. изм.", z["ed"], f["ed"]))
        if z["kolichestvo"].replace(" ", "").replace("\xa0", "") != f["kolichestvo"].replace(" ", "").replace("\xa0", ""):
            rashozhdeniya.append(Rashozhdenie(f"позиция {n}, количество", z["kolichestvo"], f["kolichestvo"]))
    return rashozhdeniya


# ------------------------------------------------------- чтение отдельных файлов

_OBRAZEC_PREDLOZHENIYA = re.compile(r"Предложение\s+\d+\s+от\s+(\d{2}\.\d{2}\.\d{4})\s*№\s*(\S+)", re.I)


def procitat_kp_iz_fajla_rascheta(put: str | Path) -> list[dict]:
    """Расчёт НМЦК «по общей сумме» в отдельном файле Word: таблица со
    строками «Предложение N от ДАТА № НОМЕР» и суммой (тот же формат, что
    в заявке — raздел izvlechenie_excel._nmck_raschet)."""
    from izvlechenie_excel import format_chislo

    doc = Document(put)
    kp = []
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells[:-1]:
                m = _OBRAZEC_PREDLOZHENIYA.search(cell.text)
                if not m:
                    continue
                summa = row.cells[-1].text.strip()
                kp.append({"vh_nomer": m.group(2), "vh_data": f"{m.group(1)}г.", "summa": format_chislo(summa)})
                break
    return kp


def procitat_instrukciyu_iz_fajla(put: str | Path, tolko_s_zagolovkom: bool = False) -> list[str]:
    """Инструкция по описанию товара в отдельном файле Word — все абзацы
    после заголовка «Инструкция по описанию товара», дословно."""
    doc = Document(put)
    abzacy = [p.text for p in doc.paragraphs]
    for i, a in enumerate(abzacy):
        if a.strip().lower().startswith("инструкция по описанию товара"):
            return abzacy[i + 1:]
    return [] if tolko_s_zagolovkom else abzacy
